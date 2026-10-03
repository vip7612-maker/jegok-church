#!/usr/bin/env python3
"""악보집 → 프레젠테이션(PPT) — 쪽마다 한 장씩 그림으로 찍어 .pptx 를 만들고, 보고 내려받는 페이지를 붙인다 (2026-10-03 교장님 지시).

  악보집 맨 위 「PPT」 단추 → 같은 폴더의 ppt.html
    · 큰 화면 + 왼쪽 작은 쪽 목록(방향키·밀기로 넘김)
    · [PPT 내려받기] → doc.pptx (A4 가로, 쪽마다 한 장)
  build.py --share 가 docsave 에 올리기 직전에 make(date, sid) 를 부른다.
  python3 accomp/pptpage.py 2026-10-04     # 따로 다시 만들 때(올리기는 build.py --share)
"""
from __future__ import annotations
import html, json, re, subprocess, sys, tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SITE = Path.home() / "dev/daily-briefing/report-site"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
W, H = 1123, 794   # A4 가로 297×210mm @96dpi


def shots(src_html: Path, n: int, out: Path) -> list[Path]:
    """악보집 HTML 의 쪽을 하나씩 찍는다(화면 단추·줄임 없이 원래 크기, 2배 선명도)."""
    out.mkdir(parents=True, exist_ok=True)
    for old in out.glob("*.jpg"): old.unlink()
    src = src_html.read_text()
    files = []
    with tempfile.TemporaryDirectory() as d:
        for i in range(1, n + 1):
            css = (f"<style>html,body{{margin:0!important;padding:0!important;background:#fff!important;overflow:hidden}}"
                   f".bar,.dlbar,#sv,#prep,#zoom,.tag{{display:none!important}}.pages{{display:block!important;padding:0!important;gap:0!important}}"
                   f".pages>.page{{display:none!important}}.pages>.page:nth-child({i}){{display:block!important;zoom:1!important;margin:0!important;box-shadow:none!important}}</style>")
            f = Path(d) / f"p{i}.html"; f.write_text(src.replace("</head>", css + "</head>", 1))
            png = Path(d) / f"p{i}.png"
            subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--window-size={W},{H}",
                            "--force-device-scale-factor=2", "--virtual-time-budget=4000", f"--screenshot={png}", f.as_uri()],
                           capture_output=True, timeout=90)
            from PIL import Image
            jpg = out / f"{i:02d}.jpg"
            Image.open(png).convert("RGB").save(jpg, "JPEG", quality=86)
            files.append(jpg)
    return files


def pptx(imgs: list[Path], dst: Path) -> None:
    from pptx import Presentation
    from pptx.util import Mm
    p = Presentation(); p.slide_width, p.slide_height = Mm(297), Mm(210)
    blank = p.slide_layouts[6]
    for im in imgs:
        s = p.slides.add_slide(blank)
        s.shapes.add_picture(str(im), 0, 0, width=p.slide_width, height=p.slide_height)
    p.save(dst)


PAGE = """<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex"><title>{title} · {label}</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#111827;color:#fff;font-family:'Apple SD Gothic Neo','Noto Sans KR',system-ui,sans-serif;height:100vh;display:flex;flex-direction:column}}
.top{{display:flex;flex-wrap:wrap;gap:8px;align-items:center;padding:8px 12px;background:#1f2937}}
.top b{{margin-right:auto;font-size:14px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;min-width:0}}
.top a{{font:700 13px inherit;text-decoration:none;border-radius:999px;padding:7px 14px;background:#fff;color:#111;white-space:nowrap}}
.top a.dl{{background:#f6c76b}}
.main{{flex:1;display:flex;min-height:0}}
#rail{{width:170px;flex:0 0 170px;overflow-y:auto;background:#0b1220;padding:8px}}
#rail img{{display:block;width:100%;margin:0 0 8px;border:2px solid transparent;border-radius:4px;cursor:pointer;background:#fff}}
#rail img.on{{border-color:#f6c76b}}
#stage{{flex:1;display:flex;align-items:center;justify-content:center;position:relative;min-width:0}}
#big{{max-width:96%;max-height:92%;background:#fff;box-shadow:0 6px 30px rgba(0,0,0,.5)}}
#pn{{position:absolute;bottom:8px;left:50%;transform:translateX(-50%);font-size:12px;color:#9ca3af}}
.nav{{position:absolute;top:0;bottom:0;width:22%;cursor:pointer}}.nav.l{{left:0}}.nav.r{{right:0}}
@media (max-width:700px){{.main{{flex-direction:column-reverse}}#rail{{width:auto;flex:0 0 86px;display:flex;gap:6px;overflow-x:auto;overflow-y:hidden}}
#rail img{{width:auto;height:100%;margin:0}}.top b{{flex:1 0 100%}}}}
</style></head><body>
<div class="top"><b>{icon} {title} · {n}{unit}</b><a class="dl" href="{dl}" download="{fname}">⬇ {label} 내려받기</a><a href="./">◀ 악보집으로</a></div>
<div class="main"><aside id="rail">{thumbs}</aside>
<div id="stage"><img id="big" alt=""><div class="nav l"></div><div class="nav r"></div><span id="pn"></span></div></div>
<script>
const S={imgs};let c=0;const big=document.getElementById('big'),pn=document.getElementById('pn'),th=[...document.querySelectorAll('#rail img')];
function go(i){{c=Math.max(0,Math.min(S.length-1,i));big.src=S[c];pn.textContent=(c+1)+' / '+S.length;th.forEach((t,k)=>t.classList.toggle('on',k===c));th[c].scrollIntoView({{block:'nearest',inline:'nearest'}});}}
th.forEach((t,k)=>t.onclick=()=>go(k));document.querySelector('.nav.l').onclick=()=>go(c-1);document.querySelector('.nav.r').onclick=()=>go(c+1);
addEventListener('keydown',e=>{{if(['ArrowRight','ArrowDown','PageDown',' '].includes(e.key)){{e.preventDefault();go(c+1)}}else if(['ArrowLeft','ArrowUp','PageUp'].includes(e.key)){{e.preventDefault();go(c-1)}}}});
let x0=null;const st=document.getElementById('stage');st.addEventListener('touchstart',e=>x0=e.touches[0].clientX,{{passive:true}});
st.addEventListener('touchend',e=>{{if(x0===null)return;const dx=e.changedTouches[0].clientX-x0;if(Math.abs(dx)>40)go(c+(dx<0?1:-1));x0=null}});
go(0);
</script></body></html>"""


def make(date: str, sid: str) -> Path:
    src = HERE / "out" / f"{date}.html"
    d = json.loads((HERE / "data" / f"{date}.json").read_text())
    n = len(re.findall(r'<section class="page" id="p\d+">', src.read_text()))
    folder = SITE / "d" / sid; folder.mkdir(parents=True, exist_ok=True)
    imgs = shots(src, n, folder / "ppt")
    fname = f"{date} 예배자 악보.pptx"
    pptx(imgs, folder / "doc.pptx")
    rel = [f"ppt/{p.name}" for p in imgs]
    (folder / "ppt.html").write_text(PAGE.format(
        title=html.escape(d["title"]), n=len(imgs), fname=html.escape(fname), icon="📽️", unit="장", label="PPT", dl="doc.pptx",
        thumbs="".join(f'<img src="{r}" alt="{i}쪽" loading="lazy">' for i, r in enumerate(rel, 1)),
        imgs=json.dumps(rel)))
    return folder


def jubo(date: str, sid: str) -> str:
    """주보 탭 — 그 주 드라이브 폴더의 「…주보.pdf」를 받아 쪽 그림 + 보기 페이지(jubo.html) + 내려받기(jubo.pdf) (2026-10-03 교장님 지시)."""
    import datetime as dt
    sys.path.insert(0, str(HERE.parent)); sys.path.insert(0, str(HERE))
    import prep, weekly, fitz
    folder = SITE / "d" / sid; folder.mkdir(parents=True, exist_ok=True)
    g = prep.G(prep.access_token()); fo = weekly.week_folder(g, dt.date.fromisoformat(date))
    pdf = None
    if fo:
        r = g.get(f"{prep.DRIVE}/files", q=f"'{fo['id']}' in parents and trashed=false and mimeType='application/pdf'",
                  fields="files(id,name,modifiedTime)", includeItemsFromAllDrives="true")
        pdf = next((f for f in sorted(r.get("files", []), key=lambda f: f["modifiedTime"], reverse=True) if "주보" in f["name"]), None)
    out = folder / "jubo"; out.mkdir(exist_ok=True)
    for old in out.glob("*.jpg"): old.unlink()
    d = json.loads((HERE / "data" / f"{date}.json").read_text())
    title = f"{int(date[5:7])}월 {int(date[8:10])}일 주일 주보"
    if not pdf:
        (folder / "jubo.html").write_text(f'<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>주보</title>'
            f'<body style="margin:0;font-family:system-ui;background:#111827;color:#fff;display:flex;flex-direction:column;align-items:center;justify-content:center;height:100vh;gap:14px">'
            f'<div style="font-size:18px">📰 {title}</div><div style="color:#9ca3af">이번 주 주보가 아직 올라오지 않았습니다.</div>'
            f'<a href="./" style="background:#fff;color:#111;border-radius:999px;padding:8px 16px;text-decoration:none;font-weight:700">◀ 악보집으로</a></body>')
        (folder / "jubo.pdf").unlink(missing_ok=True)
        return "주보 없음"
    raw = g.download(pdf["id"]); (folder / "jubo.pdf").write_bytes(raw)
    # HTML 주보 — 한글 파일 + PDF 로 원본 양식 그대로 짠 문서(jubo_form, hwp-to-html 스킬 방식)를 공유하고
    # 주보 단추(jubo.html)는 그 공유본으로 보낸다 (2026-10-03 교장님 지시)
    hwp = None
    try:
        import jubo_form
        _, url = jubo_form.make(date)
        if url:
            hwp = (f'<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
                   f'<title>{html.escape(title)}</title><meta http-equiv="refresh" content="0;url={url}">'
                   f'<script>location.replace({json.dumps(url)})</script><a href="{url}">주보 보기</a>')
    except SystemExit as e: print("주보 양식본 건너뜀:", e)
    except Exception as e: print("주보 양식본 실패:", repr(e)[:200])
    doc = fitz.open(stream=raw, filetype="pdf"); imgs = []
    for i, pg in enumerate(doc, 1):
        f = out / f"{i:02d}.jpg"; pg.get_pixmap(dpi=150).save(str(f))
        imgs.append(f)
    rel = [f"jubo/{p.name}" for p in imgs]
    if hwp:
        (folder / "jubo.html").write_text(hwp)
        return f"주보 양식본 · PDF {len(imgs)}쪽 · {pdf['name']}"
    (folder / "jubo.html").write_text(PAGE.format(
        title=html.escape(title), n=len(imgs), fname=html.escape(pdf["name"]), icon="📰", unit="쪽", label="주보", dl="jubo.pdf",
        thumbs="".join(f'<img src="{r}" alt="{i}쪽" loading="lazy">' for i, r in enumerate(rel, 1)), imgs=json.dumps(rel)))
    return f"주보 {len(imgs)}쪽 · {pdf['name']}"


if __name__ == "__main__":
    sys.path.insert(0, str(HERE))
    import publish
    sid = publish.sid_for(sys.argv[1])
    print(make(sys.argv[1], sid))
    try: print(jubo(sys.argv[1], sid))
    except Exception as e: print("주보 실패:", e)
