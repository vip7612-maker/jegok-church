#!/usr/bin/env python3
"""주일예배 PPT(프레젠테이션용) → HTML 슬라이드 (2026-10-03 교장님 지시).

PPT 를 PDF 로 바꾼 파일을 읽어 장마다
  · 배경(사진·색·악보 그림) = 글자를 지운 그림 한 장
  · 글자 = 위치·크기·색 그대로 HTML 글자(고를 수 있고 선명함)
로 다시 짠다. 다른 그림에 가려 안 보이던 글자(곡 구간 표시 등)는 빼고, 보이는 글자만 올린다.
화면: 목록(아래로 넘겨 보기) · ▶ 발표(전체 화면, 방향키·밀기) · 순서 바로가기(성경암송·찬양·설교 …) · 원본 내려받기.

  python3 accomp/slides.py <PPT.pdf> 2026-09-27 [--pptx 원본.pptx]   # 그 주 공유 폴더에 ppt.html 로
"""
from __future__ import annotations
import html, json, re, shutil, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SITE = Path.home() / "dev/daily-briefing/report-site"
SECTIONS = ["주일예배", "새벽예배", "수요예배", "금요예배", "성경암송", "찬양과경배", "사도신경", "대표기도", "교회소식", "봉헌", "성경봉독", "특송", "선교보고", "설교", "찬양과결단", "축도", "예배를마칩니다"]


def css_color(c: int) -> str:
    return f"#{c:06x}"


def parse(pdf: Path, out: Path) -> list[dict]:
    import fitz
    from PIL import Image, ImageChops
    import io
    out.mkdir(parents=True, exist_ok=True)
    for f in out.glob("*.jpg"): f.unlink()
    doc = fitz.open(pdf)
    slides = []
    for i, page in enumerate(doc):
        W, H = page.rect.width, page.rect.height
        full = Image.open(io.BytesIO(page.get_pixmap(dpi=72).tobytes("png"))).convert("RGB")
        spans = []
        for b in page.get_text("dict")["blocks"]:
            if b["type"] != 0: continue
            for ln in b["lines"]:
                for s in ln["spans"]:
                    if s["text"].strip(): spans.append(s)
        # 글자를 지운 판
        for s in spans: page.add_redact_annot(fitz.Rect(s["bbox"]), fill=None)
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
        bare = Image.open(io.BytesIO(page.get_pixmap(dpi=72).tobytes("png"))).convert("RGB")
        bare.save(out / f"{i + 1:03d}.jpg", "JPEG", quality=80)
        diff = ImageChops.difference(full, bare).convert("L")
        texts = []
        for s in spans:
            x0, y0, x1, y1 = [round(v) for v in s["bbox"]]
            box = (max(0, x0), max(0, y0), min(full.width, x1), min(full.height, y1))
            if box[2] <= box[0] or box[3] <= box[1]: continue
            if max(diff.crop(box).getextrema()) < 40:   # 지워도 안 바뀜 = 가려져 안 보이던 글자
                continue
            texts.append({"x": x0, "y": y0, "w": x1 - x0, "h": y1 - y0, "s": round(s["size"], 1),
                          "c": css_color(s["color"]), "b": "Bold" in s["font"] or bool(s["flags"] & 16), "t": s["text"]})
        texts = merge(texts, W)
        plain = " ".join(t["t"] for t in texts)
        slides.append({"n": i + 1, "w": W, "h": H, "img": f"{i + 1:03d}.jpg", "texts": texts, "plain": plain,
                       "hidden": " ".join(s["text"] for s in spans)})
    return slides


def merge(texts: list[dict], W: float) -> list[dict]:
    """같은 줄·같은 크기·같은 색 글자 조각을 한 줄로 — 낱말·문장부호 사이가 벌어지지 않게."""
    out = []
    for t in texts:
        p = out[-1] if out else None
        if (p and p["c"] == t["c"] and abs(p["s"] - t["s"]) < 1 and abs(p["y"] - t["y"]) <= max(4, t["s"] * .15)
                and 0 <= t["x"] - (p["x"] + p["w"]) <= t["s"] * 1.3):
            gap = t["x"] - (p["x"] + p["w"])
            p["t"] += (" " if gap > t["s"] * .12 and not p["t"].endswith(" ") and not t["t"].startswith(" ") else "") + t["t"]
            p["w"] = t["x"] + t["w"] - p["x"]; p["h"] = max(p["h"], t["h"]); p["b"] = p["b"] or t["b"]
        else:
            out.append(dict(t))
    for t in out:
        t["t"] = re.sub(r" {2,}", " ", t["t"]).strip()
        t["a"] = "l" if t["x"] < W * .065 else "c"   # 왼쪽 끝에서 시작하는 글(광고·사도신경)만 왼쪽 맞춤, 나머지는 원래 상자 가운데
        t["t"] = re.sub(r" ([.,)])", r"\1", re.sub(r"\( ", "(", t["t"]))
    return out


def outline(slides: list[dict], songs: list[str] | None = None) -> list[tuple[int, str]]:
    """바로가기 — 순서 표지 장과, 찬양 곡 이름(가려진 구간 표시의 첫 줄)."""
    marks, last, seen, cur = [], "", set(), ""   # cur = 바로 앞 순서 이름(곡 장은 빼고)
    def is_mark(x):   # 순서 이름(사도신경·대표기도 …)으로 시작하는 장 — 사도신경처럼 글이 길어도 표지로 본다
        f = re.sub(r"\s", "", x["plain"])
        return any(y in f[:40] for y in SECTIONS)
    for k, s in enumerate(slides):
        flat = re.sub(r"\s", "", s["plain"])
        if k == 0 and s.get("img", "").startswith("k_cover"):   # 템플릿 표지 = 그 예배 이름(관리자가 만든 예배 이름도, 2026-10-07)
            big = next((t["t"] for t in s.get("texts", []) if t.get("s", 0) >= 100 and str(t.get("t", "")).strip()), "")
            if big: marks.append((s["n"], re.sub(r"\s+", "", big))); continue
        sec = next((x for x in SECTIONS if x in flat[:40]), "")
        if sec and (len(flat) < 160 or sec == "사도신경" and "사도신경" in flat[:30] and last != "사도신경"):
            name = {"교회소식": "교회 소식", "성경봉독": "성경 봉독", "특송": "특송", "선교보고": "선교 보고", "찬양과결단": "찬양과 결단", "예배를마칩니다": "마침"}.get(sec, sec)
            # 가사가 아직 없는 틀(2026-10-03): 곡마다 「찬양과경배」 표지 장만 있다 — 같은 표지가 또 나오면 곡 자리로 보고 곡 이름을 단다
            ph = sec in ("찬양과경배", "찬양과결단") and "Praise&Worship" in flat
            empty = k + 1 >= len(slides) or is_mark(slides[k + 1])   # 바로 다음이 또 표지 = 가사 장이 없는 자리
            again = ph and sec == cur   # 바로 앞 순서도 같은 찬양 = 곡마다 다시 나오는 표지. 사이에 사도신경 등이 있으면 새 순서로 목차에 다시 단다(2026-10-07)
            if again and not empty and songs:                       # 곡마다 다시 나오는 「찬양과경배」 표지 = 다음 곡의 시작(목차도 그 곡으로)
                marks.append((s["n"], "♪ " + songs.pop(0))); last = "song"; continue
            if not again:
                marks.append((s["n"], name))
            seen.add(sec); last = sec; cur = sec
            if ph and empty and songs:
                marks.append((s["n"], "♪ " + songs.pop(0))); last = "song"
            continue
        if s.get("cover") and not sec:                            # 예배순서에서 새로 넣은 순서 표지(통성기도·간증 등) — 목차에도 (2026-10-07)
            big = next((t["t"] for t in s.get("texts", []) if t.get("s", 0) >= 100 and str(t.get("t", "")).strip()), "")
            if big: marks.append((s["n"], big.strip())); last = cur = big.strip(); continue
        if s.get("start"):                                         # add_song_frames 가 곡 첫 장에 단 곡 이름
            if songs: songs.pop(0)
            marks.append((s["n"], "♪ " + s["start"])); last = "song"; continue
        if last in ("찬양과경배", "찬양과결단"):
            m = re.match(r"\s*1\.\s*([^/0-9]+?)(?:\s+\d+\.|$)", s["hidden"])
            nm = songs.pop(0) if songs else (m.group(1).strip()[:18] if m else "찬양")
            marks.append((s["n"], "♪ " + nm)); last = "song"
    return marks


PAGE = """<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex"><title>{title}</title>
<link rel="stylesheet" href="https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable.min.css">
<style>
*{{box-sizing:border-box}}html{{-webkit-text-size-adjust:100%}}
body{{margin:0;background:#0f172a;color:#fff;font-family:'Pretendard Variable',Pretendard,'Apple SD Gothic Neo',sans-serif}}
.top{{position:sticky;top:0;z-index:20;display:flex;flex-wrap:wrap;gap:8px;align-items:center;padding:8px 12px;background:#1f2937}}
.top b{{margin-right:auto;font-size:14px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;min-width:0}}
.top a,.top button{{font:700 13px inherit;text-decoration:none;border:0;border-radius:999px;padding:7px 14px;background:#fff;color:#111;cursor:pointer;white-space:nowrap}}
.top .pr{{background:#f6c76b}}
.wrap{{display:flex}}
#toc{{position:sticky;top:var(--navh,88px);align-self:flex-start;width:190px;flex:0 0 190px;max-height:calc(100vh - var(--navh,88px));overflow-y:auto;padding:10px 8px;background:#0b1220;font-size:14.3px}}
#toc a{{display:block;color:#cbd5e1;text-decoration:none;padding:5px 8px;border-radius:6px}}#toc a:hover,#toc a.on{{background:#1e293b;color:#f6c76b}}
#toc a.song{{padding-left:18px;color:#93c5fd}}#toc a:focus,#toc a:focus-visible{{outline:none!important}}#toc a.on{{background:rgba(246,199,107,.16);color:#f6c76b;box-shadow:inset 4px 0 #f6c76b;font-weight:700}}
main{{flex:1;min-width:0;padding:14px;display:flex;flex-direction:column;align-items:center;gap:14px}}
.sl{{position:relative;width:min(100%,1100px);aspect-ratio:16/9;overflow:hidden;background:#000;box-shadow:0 4px 18px rgba(0,0,0,.4);border-radius:6px}}
.sl .in{{position:absolute;left:0;top:0;width:1440px;height:810px;transform-origin:0 0;background-size:100% 100%}}
.sl .tx{{position:absolute;white-space:pre;line-height:1;transform-origin:0 0}}
.sl .no{{position:absolute;right:8px;bottom:6px;font-size:11px;color:#fff;background:rgba(0,0,0,.45);border-radius:999px;padding:1px 7px;z-index:2}}
body.pr .top,body.pr .wsnav,body.pv .wsnav,body.pr #toc,body.pr .no{{display:none}}/* 프레젠테이션(두 번째 모니터)·전체 화면 발표엔 상단 메뉴 두 줄을 절대 안 보이게 — wsnav 의 display:block!important 를 이긴다 (2026-10-04 교장님) */body.pr .wsnav,body.pr header.wsnav,body.pr [data-wsnav]{{display:none!important}}body.pr #deck,body.pr main{{margin:0!important;padding:0!important}}body.pr main{{padding:0;gap:0}}body.pr{{background:#000;overflow:hidden}}
body.pr .sl{{display:none;position:fixed;inset:0;margin:auto;width:min(100vw,calc(100vh*16/9));border-radius:0;box-shadow:none}}body.pr .sl.cur{{display:block}}
/* 한 장 보기(기본) — 목차를 누르면 스크롤 없이 그 장만 바로 (2026-10-03 교장님 지시) */
body.one:not(.pr) main .sl{{display:none;width:min(100%,calc((100vh - var(--navh,88px) - 80px)*16/9))}}body.one:not(.pr) main .sl.cur{{display:block}}
#vbar{{display:none;align-items:center;gap:10px;font-weight:700}}body.one #vbar{{display:flex}}body.pr #vbar{{display:none}}
/* PPT 모아 보기(처음 화면, 2026-10-07 교장님): 왼쪽 목차에서 고른 순서의 모든 장을 오른쪽에 격자로. 장을 누르면 크게 */
body.grid:not(.pr) main{{display:grid;grid-template-columns:repeat(auto-fill,minmax(280px,1fr));align-content:start;align-items:start;gap:12px}}
body.grid:not(.pr) main .sl{{display:none;width:100%}}body.grid:not(.pr) main .sl.sec{{display:block;cursor:pointer}}body.grid:not(.pr) main .sl.sec:hover{{outline:3px solid #f6c76b}}
body.grid #vbar{{display:none!important}}#gbar{{grid-column:1/-1;display:none;align-items:center;gap:10px;font-weight:700;color:#cbd5e1}}body.grid:not(.pr) #gbar{{display:flex}}
#gbar b{{color:#f6c76b}}#vgrid{{display:none}}body.one #vgrid{{display:inline-block}}
#vbar button{{font:700 14px inherit;border:0;border-radius:999px;padding:8px 18px;background:#fff;color:#111;cursor:pointer}}
/* 발표자 보기(메인 모니터) — 청중 화면은 두 번째 모니터 창 (2026-10-03 교장님 지시) */
/* 3분면 (2026-10-04 교장님): 왼쪽 목차 · 가운데 이 순서의 모든 장 · 오른쪽 지금/다음 장(접었다 폈다, 가운데 막대로 폭 조절) */
#pv{{display:none;position:fixed;left:190px;top:var(--navh,88px);right:0;bottom:0;z-index:15;background:#0b1220;padding:14px;gap:0 6px;
  grid-template-columns:minmax(0,1fr) 14px var(--pvw,440px);grid-template-rows:minmax(0,1fr)}}body.pv #pv{{display:grid}}
body.pvfold #pv{{grid-template-columns:minmax(0,1fr) 0 0}}body.pvfold #pvsplit{{visibility:hidden}}
.pvt{{grid-column:1;min-width:0;min-height:0;display:flex;flex-direction:column}}
#pvsplit{{grid-column:2;cursor:col-resize;display:flex;align-items:center;justify-content:center;touch-action:none}}#pvsplit::before{{content:'';width:6px;height:120px;border-radius:99px;background:#475569}}#pvsplit:hover::before,#pvsplit.on::before{{background:#f6c76b}}
.pvside{{grid-column:3;min-width:0;min-height:0;overflow-y:auto;display:flex;flex-direction:column;gap:12px}}
#pvfold{{flex:0 0 auto;margin-left:auto;font:700 12.5px inherit;border:0;border-radius:999px;padding:5px 12px;background:#1e293b;color:#cbd5e1;cursor:pointer}}#pvfold:hover{{color:#f6c76b}}
body.pvfold .pvside{{display:none}}.pvth{{display:flex;align-items:center;gap:10px;margin:0 0 5px}}#pvscr{{flex:0 0 auto;margin-left:auto;font:800 12.5px inherit;border:0;border-radius:999px;padding:5px 12px;background:#f6c76b;color:#3b2a06;cursor:pointer}}#pvscr~#pvfold{{margin-left:0}}.pvth .pvlab{{margin:0;min-width:0}}body.pvfold #pvfold{{background:#f6c76b;color:#3b2a06}}
.pvl,.pvr{{min-width:0;display:flex;flex-direction:column;gap:12px}}
.pvlab{{font-size:12px;color:#94a3b8;margin:0 0 5px}}.pvpn{{display:flex;justify-content:space-between;align-items:flex-start}}.pvq{{width:calc((50% - 4px) * .8);cursor:pointer}}.pvq .pvlab{{margin:0 0 3px}}.pvq .sl{{border-radius:4px}}.pvq:hover .sl{{outline:2px solid #f6c76b}}.pvq .last{{font-size:11px}}.pvl .pvbox .sl{{width:100%}}.pvr .pvbox .sl{{width:100%}}
/* 발표자 보기 아래 썸네일 — 지금 노래(목차 한 구간)의 장 전부, 누르면 프레젠테이션으로 (2026-10-03 교장님 지시) */
#pvthumbs{{flex:1;min-height:0;overflow-y:auto;display:grid;grid-template-columns:repeat(auto-fill,minmax(230px,1fr));align-content:start;gap:8px;padding:2px}}
#pvthumbs .th{{min-width:0;cursor:pointer;border:3px solid transparent;border-radius:8px;padding:1px}}#pvthumbs .th:hover{{border-color:#475569}}
#pvthumbs .th.on{{border-color:#f6c76b}}#toc a.peek{{outline:2px dashed #93c5fd;outline-offset:-2px;color:#93c5fd}}#pvtlab b{{color:#93c5fd}}.pvback{{margin-left:8px;font:700 12px inherit;border:0;border-radius:999px;padding:4px 10px;background:#fff;color:#111;cursor:pointer}}#pvthumbs .sl{{width:100%;box-shadow:none;border-radius:4px}}
.pvbox .last{{aspect-ratio:16/9;display:flex;align-items:center;justify-content:center;background:#111827;border-radius:6px;color:#94a3b8}}
.pvinfo{{display:flex;justify-content:space-between;align-items:baseline;font-weight:800;font-size:26px}}#pvclock{{color:#f6c76b}}
.pvbtn{{display:flex;gap:8px}}.pvbtn button,.pvstop{{flex:1;font:700 15px inherit;border:0;border-radius:10px;padding:12px;cursor:pointer;background:#fff;color:#111}}
.pvstop{{flex:0 0 auto;background:#475569;color:#fff}}.pvnote{{font-size:12px;color:#94a3b8;line-height:1.5}}
#fstip{{position:fixed;left:50%;bottom:24px;transform:translateX(-50%);z-index:30;background:rgba(0,0,0,.7);color:#fff;padding:10px 18px;border-radius:999px;font-size:16px}}
@media (max-width:760px){{#pv{{left:0}}#toc{{display:none}}.top b{{flex:1 0 100%}}main{{padding:8px}}}}
{navcss}</style></head><body class="grid">
{nav}
<div class="wrap"><nav id="toc">{toc}</nav><main id="deck"><div id="gbar"><b id="gname"></b><span id="gcnt"></span></div>{slides}<div id="vbar"><button id="vgrid" onclick="wsGrid()">☷ 모아 보기</button><button onclick="view(cur-1)">◀ 이전</button><span id="vn"></span><button onclick="view(cur+1)">다음 ▶</button></div></main></div>
<div id="pv"><div class="pvt"><div class="pvth"><div class="pvlab" id="pvtlab">이 노래·순서의 모든 장 — 누르면 프레젠테이션에 바로 나갑니다</div><button id="pvscr" onclick="present().then(()=>{{if(scr&&!scr.closed)this.style.display='none';}})" title="두 번째 모니터에 프레젠테이션 창 열기">📺 프레젠테이션 열기</button><button id="pvfold" onclick="pvFold()" title="오른쪽 칸 접기·열기">접기 ▶</button></div><div id="pvthumbs"></div></div>
<div id="pvsplit" title="끌어서 오른쪽 폭 조절 (두 번 누르면 처음대로)"></div>
<div class="pvside"><div class="pvl"><div class="pvlab">지금 프레젠테이션</div><div id="pvcur" class="pvbox"></div><div class="pvpn"><div class="pvq" onclick="go(cur-1)" title="이전 장으로"><div class="pvlab">◀ 이전 장</div><div id="pvprev" class="pvbox"></div></div><div class="pvq" onclick="go(cur+1)" title="다음 장으로"><div class="pvlab" style="text-align:right">다음 장 ▶</div><div id="pvnext" class="pvbox"></div></div></div></div>
<div class="pvr">
<div class="pvinfo"><span id="pvn"></span><span id="pvclock">00:00</span></div>
<div class="pvbtn"><button onclick="go(cur-1)">◀ 이전</button><button onclick="go(cur+1)">다음 ▶</button></div>
<div class="pvnote">← → 방향키·스페이스로 넘김 · 왼쪽 목차를 누르면 그 장으로 · Esc 끝내기</div>
<button class="pvstop" onclick="endPV()">■ 발표 끝내기</button></div></div></div>
<script>
const D={data};
const deck=document.getElementById('deck');
function fit(){{document.querySelectorAll('.sl').forEach(sl=>{{const k=sl.clientWidth/1440;sl.querySelector('.in').style.transform='scale('+k+')';}});}}
// 글자 폭은 늘 보이는 측정 칸에서 잰다 — 발표 중 숨겨 둔 다음 장을 미리 그릴 때 폭이 0 으로 재어져
// 가운데 정렬 자막이 칸 가운데에서 시작해 잘리던 문제(2026-10-03 교장님 화면)
const MZ=document.createElement('div'); MZ.style.cssText='position:absolute;left:-99999px;top:0;visibility:hidden;white-space:pre;line-height:1'; document.body.appendChild(MZ);
function measure(t){{ MZ.style.cssText='position:absolute;left:-99999px;top:0;visibility:hidden;white-space:pre;line-height:1'+(t.css?';'+t.css:''); MZ.style.fontSize=t.s+'px'; MZ.style.fontWeight=t.b?800:500; MZ.textContent=t.t; return MZ.scrollWidth; }}
function paint(sl){{ if(sl.dataset.p) return; sl.dataset.p=1; const s=D[+sl.dataset.i]; const inn=sl.querySelector('.in');
  inn.style.backgroundImage=s.img?'url('+(/^https?:/.test(s.img)?s.img:'slides/'+s.img)+')':'none';
  let dy=0;   // 이어 붙인 문단이 PDF 보다 줄이 줄면, 그 아래 왼쪽 맞춤 글을 그만큼 올린다(빈 줄이 남지 않게)
  s.texts.forEach(t=>{{const e=document.createElement('div');e.className='tx';e.textContent=t.t;
    e.style.cssText='left:'+t.x+'px;top:'+(t.y-(t.a==='l'?dy:0))+'px;font-size:'+t.s+'px;color:'+t.c+';font-weight:'+(t.b?800:500)
      +(t.bg?';background:'+t.bg+';padding:5px 12px;border-radius:8px;line-height:1.15':'')+(t.css?';'+t.css:'');   // 작은 단추(봉독대표·회중봉독 …)
    if(t.wrap){{ e.style.width=t.w+'px'; e.style.whiteSpace=t.pre?'pre-wrap':'normal'; e.style.wordBreak='keep-all'; e.style.lineHeight=(t.lh||1.2); inn.appendChild(e);   // 원본 한 문단 = 칸 폭에서 저절로 줄바꿈
      MZ.style.cssText+=';width:'+t.w+'px;white-space:normal;word-break:keep-all;line-height:'+(t.lh||1.2); MZ.style.fontSize=t.s+'px'; MZ.style.fontWeight=t.b?800:500; MZ.textContent=t.t;
      const lines=Math.round(MZ.offsetHeight/(t.s*(t.lh||1.2))), was=Math.round(t.h/(t.s*(t.lh||1.2)));
      dy+=Math.max(0,was-lines)*t.s*(t.lh||1.2); MZ.style.cssText='position:absolute;left:-99999px;top:0;visibility:hidden;white-space:pre;line-height:1'; return; }}
    inn.appendChild(e); const w=measure(t);
    if(w>t.w&&t.w>4) e.style.transform='scaleX('+(t.w/w)+')';
    else if(t.a==='c') e.style.left=(t.x+(t.w-w)/2)+'px';}}); if(window.xpPaint) xpPaint(s,inn); }}
const io=new IntersectionObserver(es=>es.forEach(x=>{{if(x.isIntersecting)paint(x.target)}}),{{rootMargin:'800px'}});
document.querySelectorAll('.sl').forEach(sl=>io.observe(sl));
addEventListener('resize',fit); (document.fonts?document.fonts.ready:Promise.resolve()).then(fit); fit();
let cur=0; const all=[...document.querySelectorAll('.sl')];
function show(i){{cur=Math.max(0,Math.min(all.length-1,i)); all.forEach((s,k)=>s.classList.toggle('cur',k===cur)); paint(all[cur]); if(all[cur+1])paint(all[cur+1]); setTimeout(fit,0);}}
function presentHere(){{document.body.classList.add('pr'); const top=all.findIndex(s=>s.getBoundingClientRect().bottom>60); show(firstVisible());
  if(document.documentElement.requestFullscreen) document.documentElement.requestFullscreen().catch(()=>{{}});}}
function leave(){{document.body.classList.remove('pr'); all.forEach(s=>s.classList.remove('cur')); fit(); back();}}
addEventListener('keydown',e=>{{ if(!document.body.classList.contains('pr')) return;
  if(['ArrowRight','ArrowDown','PageDown',' ','Enter'].includes(e.key)){{e.preventDefault();show(cur+1)}}
  else if(['ArrowLeft','ArrowUp','PageUp','Backspace'].includes(e.key)){{e.preventDefault();show(cur-1)}}
  else if(e.key==='Home')show(0); else if(e.key==='End')show(all.length-1); else if(e.key==='Escape')leave(); }});
document.addEventListener('fullscreenchange',()=>{{ if(!SCREEN&&!document.fullscreenElement&&document.body.classList.contains('pr')) leave(); }});
deck.addEventListener('click',e=>{{ if(!document.body.classList.contains('pr')) return; show(cur+(e.clientX>innerWidth/2?1:-1)); }});
let x0=null; deck.addEventListener('touchstart',e=>x0=e.touches[0].clientX,{{passive:true}});
deck.addEventListener('touchend',e=>{{ if(x0===null||!document.body.classList.contains('pr'))return; const dx=e.changedTouches[0].clientX-x0; if(Math.abs(dx)>40)show(cur+(dx<0?1:-1)); x0=null; }});
document.querySelectorAll('#toc a:not(.bgm)').forEach(a=>a.onclick=ev=>{{ev.preventDefault(); a.blur(); if(window.bgmClose) bgmClose(); const el=document.getElementById(a.getAttribute('href').slice(1)); if(document.body.classList.contains('pv')){{peek(all.indexOf(el),a);return;}} if(document.body.classList.contains('grid')){{gridSec(a);return;}} if(ONE()){{view(all.indexOf(el));return;}} el.scrollIntoView({{behavior:'instant',block:'start'}});}});
{pvjs}
</script></body></html>"""


# 두 모니터 발표 — 메인 모니터는 발표자 보기(지금 장·다음 장·시계·목차), 두 번째 모니터는 프레젠테이션 창 (2026-10-03 교장님 지시)
# 같은 ppt.html 을 ?screen=1 로 연 창이 프레젠테이션이고, 두 창은 BroadcastChannel 로 장 번호를 주고받는다.
PVJS = r"""
const SCREEN=new URLSearchParams(location.search).has('screen');
const bc=('BroadcastChannel' in window)?new BroadcastChannel('wsppt'+location.pathname):null;
if(SCREEN){
  document.title='프레젠테이션 · '+document.title; document.body.classList.remove('one'); document.body.classList.add('pr');
  const _show=show; show=function(i,quiet){_show(i); if(!quiet&&bc)bc.postMessage({i:cur});};
  show(+(location.hash.slice(1)||0),true);
  if(bc){bc.onmessage=e=>{const m=e.data||{}; if(m.end){window.close();return;}
    if(m.bgm){bgmFS(m.bgm);return;} if(m.bgmStop){bgmFS(null);return;}
    if(typeof m.i==='number'){ if(m.i!==cur)show(m.i,true); }}; bc.postMessage({hello:1});}
  const tip=document.createElement('div'); tip.id='fstip'; tip.textContent='화면을 한 번 누르면 꽉 찬 화면이 됩니다'; document.body.appendChild(tip);
  const fs=()=>{ if(document.fullscreenElement){tip.remove();return;} document.documentElement.requestFullscreen().then(()=>tip.remove()).catch(()=>{}); };
  fs(); document.addEventListener('fullscreenchange',()=>{ if(document.fullscreenElement)tip.remove(); });
  addEventListener('click',e=>{ if(!document.fullscreenElement){e.stopImmediatePropagation(); fs();} },true);
  addEventListener('keydown',e=>{ if(e.key==='Escape'){e.stopImmediatePropagation(); return;} if(!document.fullscreenElement)fs(); },true);
}
let scr=null,t0=0,tick=null;
// 발표자 보기 위(지금·다음 장) / 아래(미리 보기) 크기 — 가운데 막대를 끌어 위아래로 조절, 이 기기에 기억 (2026-10-04 교장님)
// 3분면: 오른쪽(지금·다음 장) 폭은 가운데 막대를 좌우로 끌어 조절, 「접기 ▶」로 접었다 폈다 — 이 기기에 기억 (2026-10-04 교장님)
(function(){ const pv=document.getElementById('pv'), sp=document.getElementById('pvsplit'); if(!pv||!sp) return;
  const set=w=>{ pv.style.setProperty('--pvw',w+'px'); if(typeof fit==='function') fit(); };
  let saved=0; try{ saved=+localStorage.getItem('ws_pvw')||0; if(localStorage.getItem('ws_pvfold')==='1') document.body.classList.add('pvfold'); }catch(e){}
  const fb=()=>{ const b=document.getElementById('pvfold'); if(b) b.textContent=document.body.classList.contains('pvfold')?'열기 ◀':'접기 ▶'; }; fb();
  window.pvFold=()=>{ const on=document.body.classList.toggle('pvfold'); try{ localStorage.setItem('ws_pvfold',on?'1':'0'); }catch(e){} fb(); requestAnimationFrame(()=>{ if(typeof fit==='function') fit(); }); };
  window.pvSize=()=>{ const room=pv.clientWidth||innerWidth-190; set(Math.round(Math.max(300,Math.min(room*0.6,saved>0?saved:room*0.32)))); };
  let drag=null;
  sp.addEventListener('pointerdown',e=>{ drag={x:e.clientX,w:pv.querySelector('.pvside').getBoundingClientRect().width}; sp.classList.add('on'); sp.setPointerCapture(e.pointerId); e.preventDefault(); });
  sp.addEventListener('pointermove',e=>{ if(!drag) return; set(Math.round(Math.max(300,Math.min(pv.clientWidth*0.7,drag.w-(e.clientX-drag.x))))); });
  sp.addEventListener('pointerup',()=>{ if(!drag) return; drag=null; sp.classList.remove('on'); saved=parseInt(pv.style.getPropertyValue('--pvw'))||0; try{ localStorage.setItem('ws_pvw',saved); }catch(e){} });
  sp.addEventListener('dblclick',()=>{ saved=0; try{ localStorage.removeItem('ws_pvw'); }catch(e){} pvSize(); });
  addEventListener('resize',()=>{ if(document.body.classList.contains('pv')) pvSize(); });
})();
// BGM 을 프레젠테이션(두 번째 모니터)에서 꽉 차게 — 발표자 화면에서 고르면 BroadcastChannel 로 영상 번호가 온다 (2026-10-03 교장님)
function bgmFS(id){ let o=document.getElementById('bgmfs'); if(!id){ if(o)o.remove(); return; }
  if(!o){ o=document.createElement('div'); o.id='bgmfs'; o.style.cssText='position:fixed;inset:0;z-index:9999;background:#000'; document.body.appendChild(o); }
  o.innerHTML='<iframe src="https://www.youtube-nocookie.com/embed/'+id+'?autoplay=1&rel=0&modestbranding=1&playsinline=1" style="width:100%;height:100%;border:0" allow="autoplay; encrypted-media; fullscreen"></iframe>'; }
const ONE=()=>document.body.classList.contains('one');
function view(i){ cur=Math.max(0,Math.min(all.length-1,i)); all.forEach((s,k)=>s.classList.toggle('cur',k===cur)); paint(all[cur]); if(all[cur+1])paint(all[cur+1]); fit();
  document.getElementById('vn').textContent=(cur+1)+' / '+all.length; let on=null;
  document.querySelectorAll('#toc a').forEach(a=>{ const k=all.indexOf(document.getElementById(a.getAttribute('href').slice(1))); if(k>=0&&k<=cur)on=a; a.classList.remove('on'); });
  if(on){on.classList.add('on'); on.scrollIntoView({block:'nearest'});} window.scrollTo(0,0); }
function back(){ if(ONE()) view(cur); else all[cur].scrollIntoView({block:'center'}); }
function toggleList(){ const b=document.getElementById('lbtn');
  if(ONE()){ document.body.classList.remove('one'); all.forEach(s=>s.classList.remove('cur')); fit(); all[cur].scrollIntoView({block:'start'}); }
  else { const i=firstVisible(); document.body.classList.add('one'); view(i); }
  const o=document.getElementById('b-one'); if(o) o.classList.toggle('on',ONE()); if(b) b.classList.toggle('on',!ONE()); }
function wsView(one){ if(one!==ONE()) toggleList(); }
// ── PPT 모아 보기 (2026-10-07 교장님) — 순서 이름을 누르면 그 순서의 모든 장, ♪ 곡 이름이면 그 곡의 장만 ──
const TOC=()=>[...document.querySelectorAll('#toc a:not(.bgm)')].filter(a=>/^#s/.test(a.getAttribute('href')||''));
const idxOf=a=>all.indexOf(document.getElementById(a.getAttribute('href').slice(1)));
let gA=null;
function gridSec(a){ const L=TOC(); a=a||L[0]; if(!a) return; gA=a; const k=idxOf(a), song=a.classList.contains('song');
  const nx=L.slice(L.indexOf(a)+1).find(b=>song||!b.classList.contains('song')); const end=nx?idxOf(nx):all.length;
  document.body.classList.remove('one'); document.body.classList.add('grid');
  all.forEach((s,i)=>{ const on=i>=k&&i<end; s.classList.toggle('sec',on); s.classList.remove('cur'); if(on) paint(s); });
  L.forEach(x=>x.classList.toggle('on',x===a)); a.scrollIntoView({block:'nearest'});
  document.getElementById('gname').textContent=a.textContent.replace(/^♪\s*/,''); document.getElementById('gcnt').textContent=(end-k)+'장 · 누르면 크게';
  if(window.xpGrid) xpGrid(k,end);
  window.scrollTo(0,0); requestAnimationFrame(fit); }
window.wsGridAgain=()=>{ if(gA) gridSec(gA); };
function wsGrid(){ const L=TOC(); let a=gA; if(ONE()){ a=null; L.forEach(x=>{ if(idxOf(x)<=cur) a=x; }); } gridSec(a); }
deck.addEventListener('click',e=>{ if(!document.body.classList.contains('grid')||document.body.classList.contains('pr')) return;
  const sl=e.target.closest('.sl.sec'); if(!sl) return; document.body.classList.remove('grid'); document.body.classList.add('one'); view(all.indexOf(sl)); });
if(!SCREEN){ cur=0; requestAnimationFrame(()=>{ if(!document.body.classList.contains('pv')) gridSec(); });
  addEventListener('keydown',e=>{ const c=document.body.classList; if(!ONE()||c.contains('pr')||c.contains('pv')) return;
    if(['ArrowRight','ArrowDown','PageDown',' '].includes(e.key)){e.preventDefault();view(cur+1)}
    else if(['ArrowLeft','ArrowUp','PageUp'].includes(e.key)){e.preventDefault();view(cur-1)}
    else if(e.key==='Home')view(0); else if(e.key==='End')view(all.length-1); });
  deck.addEventListener('touchend',e=>{ if(x0===null||!ONE()||document.body.classList.contains('pr'))return; const dx=e.changedTouches[0].clientX-x0; if(Math.abs(dx)>40)view(cur+(dx<0?1:-1)); x0=null; }); }
function firstVisible(){if(ONE())return cur; const i=all.findIndex(s=>s.getBoundingClientRect().bottom>60); return i<0?0:i;}
async function present(){
  if(SCREEN) return;
  let other=null;
  try{ if('getScreenDetails' in window){ const sd=await getScreenDetails(); other=sd.screens.find(s=>s!==sd.currentScreen)||null; } }catch(e){}
  window.WS_OTHER=other;   // BGM 창도 같은 두 번째 모니터에 띄우려고 기억해 둔다
  if(!other&&window.screen.isExtended===false){ presentHere(); return; }   // 모니터가 하나뿐이면 예전처럼 이 창에서 전체 화면
  const start=firstVisible();
  const f=other?'popup,left='+other.availLeft+',top='+other.availTop+',width='+other.availWidth+',height='+other.availHeight+',fullscreen'
               :'popup,width=1280,height=720';
  scr=window.open(location.pathname+'?screen=1#'+start,'wsppt_screen',f);
  if(!scr){ alert('프레젠테이션 창이 막혔습니다. 주소창 오른쪽에서 팝업·창 관리를 「허용」한 뒤 ▶ 예배용을 다시 눌러 주세요.'); return; }
  if(!other) alert('프레젠테이션 창을 두 번째 모니터로 끌어 놓고, 그 창을 한 번 누르면 꽉 찬 화면이 됩니다.');
  enterPV(start);
}
function clock(){const s=Math.floor((Date.now()-t0)/1000); document.getElementById('pvclock').textContent=String(Math.floor(s/60)).padStart(2,'0')+':'+String(s%60).padStart(2,'0');}
function enterPV(i){ document.body.classList.add('pv'); if(window.pvSize) pvSize(); t0=Date.now(); clearInterval(tick); tick=setInterval(clock,1000); clock(); go(i); }
function mount(id,i){ const box=document.getElementById(id); box.replaceChildren();
  if(i>=all.length){ box.innerHTML='<div class="last">마지막 장입니다</div>'; return; }
  if(i<0){ box.innerHTML='<div class="last">첫 장입니다</div>'; return; }
  paint(all[i]); const c=all[i].cloneNode(true); c.removeAttribute('id'); c.classList.remove('cur'); box.appendChild(c); }
function marks(){ return [...document.querySelectorAll('#toc a')].map(a=>all.indexOf(document.getElementById(a.getAttribute('href').slice(1)))).filter(k=>k>=0).sort((a,b)=>a-b); }
function songRange(i){ let s=0,e=all.length; for(const k of marks()){ if(k<=i) s=k; else { e=k; break; } } return [s,e]; }   // 목차 표시 사이 = 한 노래(순서)
let thR=null;
// 발표 중 목차를 누르면 그 순서의 장들만 아래에 펼쳐 미리 본다 — 프레젠테이션은 장(썸네일)을 눌러야 넘어간다 (2026-10-04 교장님: 다음 곡 준비)
let peekAt=null;
function peek(k,a){ if(k<0) return; peekAt=k; thumbs(k); document.querySelectorAll('#toc a').forEach(x=>x.classList.toggle('peek',x===a&&!x.classList.contains('on')));
  const [s]=songRange(k), lab=document.getElementById('pvtlab'); const same=songRange(cur)[0]===s;
  lab.innerHTML=same?'이 노래·순서의 모든 장 — 누르면 프레젠테이션에 바로 나갑니다'
    :'👀 미리 보기: <b>'+(a?a.textContent:'')+'</b> — 프레젠테이션은 그대로입니다. 장을 눌러야 넘어갑니다 <button class="pvback" onclick="unpeek()">↩ 지금 순서로</button>';
  document.getElementById('pvthumbs').scrollTop=0; }
function unpeek(){ peekAt=null; document.querySelectorAll('#toc a.peek').forEach(x=>x.classList.remove('peek')); document.getElementById('pvtlab').textContent='이 노래·순서의 모든 장 — 누르면 프레젠테이션에 바로 나갑니다'; thumbs(); }
function thumbs(at){ const [s,e]=songRange(at==null?cur:at), box=document.getElementById('pvthumbs');
  if(!thR||thR[0]!==s||thR[1]!==e){ box.replaceChildren();
    for(let k=s;k<e;k++){ paint(all[k]); const c=all[k].cloneNode(true); c.removeAttribute('id'); c.classList.remove('cur');
      const w=document.createElement('div'); w.className='th'; w.dataset.k=k; w.title=(k+1)+'번째 장 — 누르면 프레젠테이션으로'; w.appendChild(c); w.onclick=()=>go(k); box.appendChild(w); }
    if(window.xpThumbs) xpThumbs(box,s,e); thR=[s,e]; }
  box.querySelectorAll('.th').forEach(t=>t.classList.toggle('on',+t.dataset.k===cur)); fit();
  const on=box.querySelector('.th.on'); if(on) on.scrollIntoView({block:'nearest'}); }
function go(i,quiet){ const fe=document.activeElement; if(fe&&fe.closest&&fe.closest('#toc')) fe.blur();   // 목차 초점 테두리가 지난 항목에 남지 않게
  cur=Math.max(0,Math.min(all.length-1,i)); mount('pvcur',cur); mount('pvprev',cur-1); mount('pvnext',cur+1); if(peekAt!==null) unpeek(); else thumbs(); fit();
  document.getElementById('pvn').textContent=(cur+1)+' / '+all.length;
  let on=null;
  document.querySelectorAll('#toc a').forEach(a=>{ const k=all.indexOf(document.getElementById(a.getAttribute('href').slice(1))); if(k>=0&&k<=cur)on=a; a.classList.remove('on'); });
  if(on){on.classList.add('on'); on.scrollIntoView({block:'nearest'});}
  if(!quiet&&bc) bc.postMessage({i:cur}); }
function endPV(){ document.body.classList.remove('pv'); clearInterval(tick); if(bc)bc.postMessage({end:1}); try{scr&&scr.close()}catch(e){} scr=null; back(); }
if(!SCREEN&&bc) bc.onmessage=e=>{ const m=e.data||{}; if(!document.body.classList.contains('pv'))return;
  if(m.hello) bc.postMessage({i:cur}); else if(typeof m.i==='number'&&m.i!==cur) go(m.i,true); };
addEventListener('keydown',e=>{ if(!document.body.classList.contains('pv')) return;
  if(['ArrowRight','ArrowDown','PageDown',' ','Enter'].includes(e.key)){e.preventDefault();go(cur+1)}
  else if(['ArrowLeft','ArrowUp','PageUp','Backspace'].includes(e.key)){e.preventDefault();go(cur-1)}
  else if(e.key==='Home')go(0); else if(e.key==='End')go(all.length-1); else if(e.key==='Escape')endPV(); });

// ── BGM·유튜브는 늘 따로 새 창(jegok_worship/bgm.html) — 슬라이드를 넘겨도 끊기지 않는다 (2026-10-04 교장님) ──
// BGM 창 주소: report-site 에서는 /jegok_worship/bgm.html, worship-desk 에서는 /<교회>/bgm.html — 예전 주소만 써서 worship-desk 에서 404 (2026-10-08)
window.wsBgmUrl=()=>/(^|\.)report-site|^localhost|^127\./.test(location.hostname)?'/jegok_worship/bgm.html':'/'+(location.pathname.split('/')[1]||'jegok')+'/bgm.html';
window.wsBgm={ w:null, id:null,
  open(id){ if(!id) return null; if(this.w&&!this.w.closed&&this.id===id){ try{this.w.focus();}catch(e){} return this.w; }
    const o=window.WS_OTHER, f=o?('popup,left='+o.availLeft+',top='+o.availTop+',width='+o.availWidth+',height='+o.availHeight):'popup,width=1280,height=720';
    const u=wsBgmUrl()+'?v='+encodeURIComponent(id);
    if(this.w&&!this.w.closed){ try{ this.w.location.href=u; this.id=id; this.w.focus(); wsBgmBtn(); return this.w; }catch(e){} }
    this.w=window.open(u,'wsbgm',f); this.id=this.w?id:null;
    if(!this.w) alert('BGM 창이 막혔습니다. 주소창 오른쪽에서 팝업을 「허용」한 뒤 다시 눌러 주세요.');
    wsBgmBtn(); return this.w; },
  stop(){ try{ if(this.w&&!this.w.closed) this.w.close(); }catch(e){} if('BroadcastChannel' in window) new BroadcastChannel('wsbgm').postMessage({stop:1}); this.w=null; this.id=null; wsBgmBtn(); },
  on(){ return !!(this.w&&!this.w.closed); } };
function wsBgmBtn(){ if(SCREEN) return; let b=document.getElementById('bgmstopf');
  if(!b){ b=document.createElement('button'); b.id='bgmstopf'; b.textContent='■ BGM 멈춤';
    b.style.cssText="position:fixed;left:18px;bottom:18px;z-index:65;display:none;font:800 15px 'Pretendard Variable',sans-serif;border:0;border-radius:999px;padding:11px 18px;background:#dc2626;color:#fff;box-shadow:0 4px 16px rgba(0,0,0,.4);cursor:pointer";
    b.onclick=()=>{ wsBgm.stop(); if(typeof bgmStop==='function') try{bgmStop();}catch(e){} }; document.body.appendChild(b); }
  b.style.display=wsBgm.on()?'block':'none'; }
if(!SCREEN) setInterval(wsBgmBtn,1000);

// ── 교회 소식 급히 고치기 (2026-10-04 교장님) ─────────────────────────────
// 「교회 소식」 표지 다음부터 「봉헌」 표지 앞까지가 소식 장. ✏️ 를 누르면 줄마다 고칠 수 있고,
// 저장하면 프레젠테이션(두 번째 모니터)에도 바로 바뀌고 서버에 남아 새로 열어도 유지된다.
const NEWS=(()=>{ const T=i=>D[i].texts.map(t=>t.t).join(' '); let s=-1,e=-1;
  for(let i=0;i<D.length;i++){ const t=T(i).replace(/\s/g,''); if(s<0&&/교회소식|Announcements/.test(t)) s=i; else if(s>=0&&/^(Offering|봉헌)|Offering봉헌/.test(t)){ e=i; break; } }
  const r=[]; if(s>=0) for(let i=s+1;i<(e<0?Math.min(D.length,s+8):e);i++) r.push(i); return r; })();
// 순서 장(대표기도·봉헌·성경봉독·특송·설교·선교보고…) — 맡은 사람 이름이 현장에서 바뀌면 그 자리에서 고친다 (2026-10-04 교장님)
const ROLE=D.map((s,i)=>i).filter(i=>{ const t=D[i].texts.map(x=>x.t).join(' ');
  return D[i].texts.length<=8 && /^(Prayer|Offering|Scripture\s*Reading|Special\s*Praise|Sermon|Mission\s*Report|Benediction|대표\s*기도|봉\s*헌|성경\s*봉독|특\s*송|설\s*교|선교\s*보고)/.test(t.trim()); });
const EDIT=[...new Set([...NEWS,...ROLE])];
// 대표기도·봉헌·특송은 맡은 분 이름 칸만 고친다(제목·안내 문구는 그대로) — 2026-10-04 교장님
const NAMEONLY=i=>/^(Prayer|Offering|Special\s*Praise|대표\s*기도|봉\s*헌|특\s*송)/.test(D[i].texts.map(x=>x.t).join(' ').trim());
const ISNAME=t=>t.s>=40&&t.s<100;
const ORIG=D.map(s=>s.texts.map(t=>t.t).join('\n'));
function repaint(i){ const sl=all.find(x=>+x.dataset.i===i); if(!sl) return; sl.querySelectorAll('.in .tx').forEach(x=>x.remove()); delete sl.dataset.p; paint(sl);
  if(document.body.classList.contains('pv')){ thR=null; go(cur,true); } fit(); }
function applyNews(i,texts){ if(!D[i]) return; texts.forEach((t,j)=>{ if(D[i].texts[j]) D[i].texts[j].t=t; }); repaint(i); }
if(bc){ const prev=bc.onmessage; bc.onmessage=e=>{ const m=e.data||{}; if(m.news){ applyNews(m.news.i,m.news.texts); return; } if(prev) prev(e); }; }
// 고친 글자(교회 소식·이름) — 열 때, 그리고 30초마다(예배준비 화면에서 이름을 고치면 띄워 둔 PPT 에도, 2026-10-07)
let lastNews='';
function loadNews(){ fetch('/api/worship?news='+NDATE,{cache:'no-store'}).then(r=>r.text()).then(t=>{ if(t===lastNews) return; lastNews=t;
  (JSON.parse(t)||[]).forEach(r=>{ if(r.orig===ORIG[r.i]) applyNews(r.i,r.texts); }); }).catch(()=>{}); }
loadNews(); setInterval(loadNews,30000);
if(!SCREEN&&EDIT.length){
  const st=document.createElement('style'); st.textContent=`#nbtn{position:fixed;right:18px;bottom:18px;z-index:60;display:none;font:800 16px 'Pretendard Variable',sans-serif;border:0;border-radius:999px;padding:12px 20px;background:#f6c76b;color:#3b2a06;box-shadow:0 4px 16px rgba(0,0,0,.4);cursor:pointer}
#nbox{position:fixed;inset:0;z-index:70;display:none;background:rgba(0,0,0,.6);align-items:center;justify-content:center}
#nbox .p{background:#fff;color:#111;border-radius:14px;width:min(860px,94vw);max-height:90vh;overflow:auto;padding:18px 20px;font:15px 'Pretendard Variable',sans-serif}
#nbox h3{margin:0 0 6px}#nbox .h{color:#64748b;font-size:13px;margin-bottom:10px}
#nbox textarea{width:100%;box-sizing:border-box;font:16px 'Pretendard Variable',sans-serif;border:1px solid #cbd5e1;border-radius:8px;padding:8px;margin:4px 0;resize:vertical}
#nbox .b{display:flex;gap:8px;justify-content:flex-end;margin-top:10px}#nbox .b button{font:700 15px inherit;border:0;border-radius:999px;padding:10px 18px;cursor:pointer}
#nbox .ok{background:#1f2937;color:#fff}#nbox .no{background:#e2e8f0}#nbox .m{font-size:13px;color:#b45309;margin-right:auto;align-self:center}`;
  document.head.appendChild(st);
  const btn=document.createElement('button'); btn.id='nbtn'; btn.textContent='✏️ 교회 소식 수정'; document.body.appendChild(btn);
  const box=document.createElement('div'); box.id='nbox'; box.innerHTML='<div class="p"><h3>✏️ 교회 소식 수정</h3><div class="h"></div><div class="f"></div><div class="b"><span class="m"></span><button class="no">닫기</button><button class="ok">저장 — 프레젠테이션에 바로 반영</button></div></div>'; document.body.appendChild(box);
  const here=()=>{ const k=document.body.classList.contains('pv')||ONE()?cur:firstVisible(); return all[k]?+all[k].dataset.i:k; };
  const sync=()=>{ const h=here(); btn.style.display=EDIT.includes(h)?'block':'none'; btn.textContent=NEWS.includes(h)?'✏️ 교회 소식 수정':NAMEONLY(h)?'✏️ 이름 수정':'✏️ 이름·내용 수정'; };
  const _go=go; go=function(i,q){ _go(i,q); sync(); }; const _view=view; view=function(i){ _view(i); sync(); };
  addEventListener('scroll',()=>{ if(!ONE()) sync(); },{passive:true}); setInterval(sync,800); sync();
  let ei=-1;
  btn.onclick=()=>{ ei=here(); const f=box.querySelector('.f'); f.replaceChildren();
    const only=!NEWS.includes(ei)&&NAMEONLY(ei); box.querySelector('h3').textContent=NEWS.includes(ei)?'✏️ 교회 소식 수정':only?'✏️ 맡은 분 이름 수정 (오늘만)':'✏️ 이름·내용 수정 (오늘만)'; box.querySelector('.h').textContent=(ei+1)+'번째 장 · 줄마다 고친 뒤 저장하세요 (이번 주 PPT 에만 반영, 주보는 그대로)';
    D[ei].texts.forEach((t,j)=>{ if(only&&!ISNAME(t)) return; const a=document.createElement('textarea'); a.dataset.j=j; a.value=t.t; a.rows=Math.max(1,Math.ceil(t.t.length/48)); if(only) a.placeholder='맡은 분 이름 (예: 홍길동 장로)'; f.appendChild(a); });
    box.querySelector('.m').textContent=''; box.style.display='flex'; };
  box.addEventListener('keydown',e=>e.stopPropagation(),true);
  box.querySelector('.no').onclick=()=>{ box.style.display='none'; };
  box.querySelector('.ok').onclick=()=>{ const texts=D[ei].texts.map(t=>t.t); box.querySelectorAll('textarea').forEach(a=>{ texts[+a.dataset.j]=a.value; }); const m=box.querySelector('.m');
    applyNews(ei,texts); if(bc) bc.postMessage({news:{i:ei,texts}});
    m.textContent='저장 중…';
    fetch('/api/worship',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({action:'news',date:NDATE,key:NKEY,i:ei,orig:ORIG[ei],texts})})
      .then(r=>r.json()).then(j=>{ if(j.ok){ m.textContent='✅ 저장했습니다'; setTimeout(()=>{box.style.display='none';},700); } else m.textContent='화면엔 바꿨지만 저장 실패: '+(j.error||''); })
      .catch(()=>{ m.textContent='화면엔 바꿨지만 저장 실패 — 인터넷 연결을 확인해 주세요'; }); };
}
"""


# ── BGM — 예배를 마친 뒤 성도의 교제 때 틀 신나는 CCM (2026-10-03 교장님) ─────────────────
# ppt.html 을 만들 때 유튜브에서 찾아 퍼가기(임베드)가 되는 것만 5개. 같은 주는 out/bgm/<날짜>.json 에 두어 다시 만들어도 그대로.
BGM_Q = ["신나는 CCM 찬양 모음", "기쁨의 찬양 CCM 플레이리스트", "경쾌한 CCM 찬양 연속듣기", "밝은 CCM 예배찬양 플레이리스트",
         "신나는 찬양 메들리", "드라이브 CCM 찬양 플레이리스트", "업비트 CCM 찬양 모음"]
# BGM 고르기 단추 — 단추마다 유튜브 5개 (2026-10-07 교장님: 새벽예배든 다른 예배든)
BGM_CATS = {
    # 기도용은 노래가 있는 CCM 이 늘 섞이게 — 찾는 말 두 개 중 하나는 꼭 「CCM」 (2026-10-08 국장님: 기도할 때 트는 CCM 도 목록에)
    "기도용 CCM(경건)": ["기도할 때 듣는 CCM 찬양 플레이리스트", "기도할 때 듣는 찬양 피아노", "잔잔한 기도 CCM 찬양 모음", "잔잔한 묵상 찬양 피아노 연주",
                       "묵상 CCM 찬양 연속듣기", "경건한 찬송가 피아노 연주"],
    "기도용 CCM(열정)": ["기도할 때 듣는 CCM 찬양 연속듣기", "뜨거운 기도 찬양 연속듣기", "기도회 CCM 찬양 플레이리스트", "부흥 찬양 기도 모음",
                       "성령 CCM 찬양 기도회", "성령 찬양 기도회 연속듣기"],
    # 기도회 때 깔아 두는 반주(노래 없는 연주·패드) — 이 칸만 「반주·MR」 제목도 받는다 (2026-10-08 국장님)
    "기도회 반주": ["기도회 반주 음악", "통성기도 반주 음악", "기도회 배경음악 피아노 패드", "중보기도 반주 연속",
                 "기도 반주 키보드 패드", "기도회 BGM 반주"],
    "경쾌한 CCM": BGM_Q,
    "JAZZ CCM": ["재즈 찬양 피아노", "CCM 재즈 플레이리스트", "재즈 찬송가 연주", "jazz worship piano playlist"],
}


def _yt(q: str, accomp: bool = False) -> list[dict]:
    import subprocess
    try:
        r = subprocess.run(["yt-dlp", "--no-warnings", "--skip-download", "--print",
                            "%(id)s\t%(title)s\t%(duration)s\t%(playable_in_embed)s\t%(channel)s", f"ytsearch8:{q}"],
                           capture_output=True, text=True, timeout=180)
    except Exception:
        return []
    out = []
    for ln in r.stdout.splitlines():
        p = ln.split("\t")
        if len(p) < 5 or p[3] != "True" or not p[2].isdigit() or int(p[2]) < 180: continue
        if re.search(r"(?i)shorts", p[1]) or (not accomp and re.search(r"(?i)MR|반주|inst", p[1])): continue
        out.append({"id": p[0], "title": p[1], "sec": int(p[2]), "ch": p[4]})
    return out


def _pick(cands: list[dict], n: int) -> list[dict]:
    seen, out, extra = set(), [], []
    for v in cands:
        if v["id"] in seen: continue
        seen.add(v["id"]); (out if v["ch"] not in {x["ch"] for x in out} else extra).append(v)   # 채널이 겹치지 않게 먼저
    return (out + extra)[:n]


def bgm(date: str, n: int = 5) -> dict[str, list[dict]]:
    """단추별 BGM — {단추 이름: [유튜브 5개]}. 그 때(날짜·예배)마다 한 번 찾아 out/bgm/<때>.json 에 둔다. 주마다 찾는 말을 돌린다."""
    import datetime as dt, svc
    f = HERE / "out" / "bgm" / f"{date}.json"
    old = json.loads(f.read_text()) if f.exists() else {}
    if isinstance(old, list): old = {"경쾌한 CCM": old}          # 예전 꼴(목록 하나) = 경쾌한 CCM
    wk = dt.date.fromisoformat(svc.ymd(date)).isocalendar()[1]
    cats = {}
    for name, qs in BGM_CATS.items():
        if old.get(name): cats[name] = old[name]; continue
        got = []
        two = (qs[wk % len(qs)], qs[(wk + 3) % len(qs)])
        if name.startswith("기도용"):                       # 짝수 칸 = CCM, 홀수 칸 = 피아노·기도회 → 한 주에 둘 다
            k = (wk % (len(qs) // 2)) * 2; two = (qs[k], qs[k + 1])
        for q in two: got += _yt(q, accomp=name == "기도회 반주")
        cats[name] = _pick(got, n)
    if any(cats.values()):
        f.parent.mkdir(parents=True, exist_ok=True); f.write_text(json.dumps(cats, ensure_ascii=False, indent=1))
    return cats


def bgm_html(cats, first: str = "") -> str:
    if isinstance(cats, list): cats = {"경쾌한 CCM": cats}
    cats = {k: v for k, v in (cats or {}).items() if v}
    if not cats: return ""
    first = first if first in cats else next(iter(cats))
    tm = lambda s: f"{s // 3600}시간 {s % 3600 // 60}분" if s >= 3600 else f"{s // 60}분"
    card = lambda v: (f'<button class="bgmc" data-id="{html.escape(v["id"])}"><img src="https://i.ytimg.com/vi/{html.escape(v["id"])}/mqdefault.jpg" alt="" loading="lazy">'
                      f'<span>{html.escape(v["title"])}</span><small>{html.escape(v["ch"])} · {tm(v["sec"])}</small></button>')
    tabs = "".join(f'<button type="button" class="bgmt{" on" if k == first else ""}" data-cat="{html.escape(k)}">{html.escape(k)}</button>' for k in cats)
    lists = "".join(f'<div class="bgml{" on" if k == first else ""}" data-cat="{html.escape(k)}">' + "".join(card(v) for v in L) + '</div>' for k, L in cats.items())
    return ('<div id="bgmbox"><div class="bgmh"><b>🎵 BGM</b><span id="bgmnow"></span><button class="bgmx" onclick="bgmClose()">닫기 (음악은 계속)</button></div>'
            '<div class="bgmtabs">' + tabs + '</div><div id="bgmmode"></div><div id="bgmplay"></div>' + lists + '<p class="bgmnote">누르면 BGM 전용 새 창(두 번째 모니터)에서 나옵니다 · 슬라이드를 넘기거나 다른 순서로 가도 음악은 계속 · 끝낼 때 ■ BGM 멈춤</p></div>'
            """<script>
// 🎵 BGM 단추·목차 → 슬라이드 화면과 따로, 두 번째 모니터에 BGM 전용 창(목록 + 재생)을 연다 (2026-10-04 교장님)
function bgmOpen(){ const items=[...document.querySelectorAll('#bgmbox .bgml.on .bgmc')].map(c=>({id:c.dataset.id,title:c.querySelector('span').textContent,ch:(c.querySelector('small').textContent.split(' · ')[0]||'')}));
  const o=window.WS_OTHER, f=o?('popup,left='+o.availLeft+',top='+o.availTop+',width='+o.availWidth+',height='+o.availHeight):'popup,width=1280,height=720';
  const w=window.open(wsBgmUrl()+'#'+encodeURIComponent(JSON.stringify(items)),'wsbgm',f);
  if(w){ wsBgm.w=w; wsBgm.id='list'; try{w.focus();}catch(e){} wsBgmBtn(); return; }
  bgmMode(); document.body.classList.add('bgm'); document.querySelectorAll('#toc a').forEach(a=>a.classList.toggle('on',a.classList.contains('bgm'))); }   // 팝업이 막혔을 때만 예전처럼 화면 안에
function bgmScreen(){ return typeof scr!=='undefined' && scr && !scr.closed && typeof bc!=='undefined' && bc; }
// BGM 은 따로 새 창(jegok_worship/bgm.html)으로 — 슬라이드를 넘겨도 끊기지 않는다 (2026-10-04 교장님)
function bgmWin(id){ return wsBgm.open(id); }
function bgmStop(){ if(wsBgm.on()) wsBgm.stop();
  if(bgmScreen()) bc.postMessage({bgmStop:1}); document.getElementById('bgmplay').innerHTML=''; document.getElementById('bgmnow').textContent=''; document.querySelectorAll('.bgmc').forEach(x=>x.classList.remove('on')); }
function bgmMode(){ const m=document.getElementById('bgmmode'); if(!m) return; const on=bgmScreen();
  m.innerHTML=window.WS_OTHER?'📺 고르면 두 번째 모니터에 BGM 전용 창이 뜹니다 — 슬라이드와 따로 돌아갑니다'
               :'고르면 BGM 전용 새 창이 뜹니다 — 두 번째 모니터로 끌어 놓고 한 번 누르면 전체 화면 (▶ 예배용을 먼저 켜면 자리를 알아서 잡습니다)'; }
async function bgmOpenScreen(){ if(typeof present==='function'){ await present(); bgmOpen(); } }
function bgmClose(){ document.body.classList.remove('bgm'); const b=document.querySelector('#toc a.bgm'); if(b) b.classList.remove('on'); }
addEventListener('DOMContentLoaded',()=>{
  const t=document.querySelector('#toc a.bgm'); if(t) t.onclick=e=>{ e.preventDefault(); t.blur(); bgmOpen(); };
  // 단추(기도용 CCM 경건·열정 / 경쾌한 CCM / JAZZ CCM)를 누르면 그 5개가 보인다
  document.querySelectorAll('#bgmbox .bgmt').forEach(b=>b.onclick=()=>{ const k=b.dataset.cat;
    document.querySelectorAll('#bgmbox .bgmt').forEach(x=>x.classList.toggle('on',x===b));
    document.querySelectorAll('#bgmbox .bgml').forEach(x=>x.classList.toggle('on',x.dataset.cat===k)); });
  document.getElementById('bgmbox').onclick=e=>{ const c=e.target.closest('.bgmc'); if(!c) return;
    document.querySelectorAll('.bgmc').forEach(x=>x.classList.toggle('on',x===c));
    const name=c.querySelector('span').textContent.slice(0,40), P=document.getElementById('bgmplay');
    if(bgmScreen()) bc.postMessage({bgmStop:1});      // 예전 방식(프레젠테이션 위 덮개)이 떠 있으면 걷는다
    if(!bgmWin(c.dataset.id)) return;
    P.innerHTML='<div class="bgmon">🎵 BGM 전용 창에서 재생 중 — 다른 슬라이드로 가도 끊기지 않습니다 · 그 창을 한 번 누르면 전체 화면 <button class="bgmx" onclick="bgmStop()">■ BGM 멈춤</button></div>';
    document.getElementById('bgmnow').textContent='BGM 창 재생: '+name; };
  bgmMode();
});
</script>""")


BGM_CSS = """
#bgmbox{display:none;position:fixed;top:var(--navh,88px);left:190px;right:0;bottom:0;z-index:25;background:#0f172a;overflow:auto;padding:16px 20px}
body.bgm #bgmbox{display:block}body.pv #bgmbox{z-index:60;left:0}body.pr #bgmbox{display:none!important}
.bgmh{display:flex;align-items:center;gap:12px;margin-bottom:12px}.bgmh b{font-size:18px}#bgmnow{flex:1;color:#f6c76b;font-size:13px;overflow:hidden;white-space:nowrap;text-overflow:ellipsis}
.bgmx{font:700 13px inherit;border:0;border-radius:999px;padding:7px 14px;background:#fff;color:#111;cursor:pointer}
#bgmmode{color:#cbd5e1;font-size:13.5px;margin:-4px 0 12px}
.bgmtabs{display:flex;flex-wrap:wrap;gap:8px;margin:0 0 12px}.bgmt{font:700 14px inherit;border:1px solid #334155;border-radius:999px;padding:8px 16px;background:#1e293b;color:#cbd5e1;cursor:pointer}
.bgmt.on{background:#f6c76b;border-color:#f6c76b;color:#0b1430}#bgmbox .bgml:not(.on){display:none}#bgmmode .bgmx{margin-left:8px}
.bgmon{display:flex;align-items:center;gap:12px;justify-content:center;background:#1e293b;border:1px solid #f6c76b;border-radius:10px;padding:18px;margin:0 auto 14px;max-width:960px;font-weight:700;color:#f6c76b}
#bgmplay iframe{width:min(100%,960px);aspect-ratio:16/9;border:0;border-radius:10px;display:block;margin:0 auto 14px;background:#000}
.bgml{display:grid;grid-template-columns:repeat(auto-fill,minmax(220px,1fr));gap:12px}
.bgmc{display:flex;flex-direction:column;gap:6px;text-align:left;border:2px solid transparent;border-radius:10px;background:#1e293b;color:#e2e8f0;padding:6px;cursor:pointer;font:600 13.5px inherit}
.bgmc:hover{border-color:#475569}.bgmc.on{border-color:#f6c76b}.bgmc img{width:100%;aspect-ratio:16/9;object-fit:cover;border-radius:6px}
.bgmc span{display:-webkit-box;-webkit-line-clamp:2;-webkit-box-orient:vertical;overflow:hidden}.bgmc small{color:#94a3b8;font-weight:400}
.bgmnote{color:#94a3b8;font-size:12.5px}#toc a.bgm{margin-top:8px;border-top:1px solid #1e293b;color:#f6c76b}#toc a.xmadd{margin-top:8px;border-top:1px solid #1e293b;color:#93c5fd}#toc a.xmenu{color:#e2e8f0}
@media (max-width:760px){#bgmbox{left:0}}
"""

# ── 💬 예배팀 소통 — PPT 왼쪽 목차 아래 채팅. 체크 안 한 메시지는 상단 메뉴 전광판에서 깜빡이며 번갈아 (2026-10-04 교장님) ──
CHAT_HTML = ('<div id="chat"><div class="chh">💬 예배팀 소통</div>'
             '<div class="chin"><textarea id="chtx" rows="2" placeholder="메시지 쓰기 (Enter 발송 · Shift+Enter 줄바꿈)"></textarea>'
             '<button type="button" id="chgo">챗발송</button></div><div id="chlist"></div></div>')

CHAT_CSS = """
#chat{margin-top:10px;border-top:1px solid #1e293b;padding-top:10px}
#chat .chh{font-weight:800;color:#e2e8f0;padding:0 4px 6px}
#chat .chin{display:flex;flex-direction:column;gap:6px}
#chtx{width:100%;box-sizing:border-box;resize:vertical;min-height:52px;font:14px 'Pretendard Variable',sans-serif;border:1px solid #334155;border-radius:8px;background:#0f172a;color:#fff;padding:7px 8px}
#chgo{font:800 14px 'Pretendard Variable',sans-serif;border:0;border-radius:8px;padding:8px;background:#f6c76b;color:#3b2a06;cursor:pointer}
#chlist{margin-top:8px;max-height:186px;overflow-y:auto;display:flex;flex-direction:column;gap:6px}
#chlist .cm{display:flex;gap:7px;align-items:flex-start;background:#1e293b;border-radius:8px;padding:7px 8px;font-size:13.5px;line-height:1.35;color:#f1f5f9;word-break:keep-all;overflow-wrap:anywhere}
#chlist .cm input{margin:2px 0 0;width:18px;height:18px;flex:0 0 auto;cursor:pointer;accent-color:#16a34a}
#chlist .cm small{display:block;color:#94a3b8;font-size:11px;margin-top:2px}
#chlist .none{color:#64748b;font-size:12.5px;padding:2px 4px}
body.pr #chat{display:none}
"""

CHAT_JS = r"""
(function(){
  const send=async()=>{ const ta=document.getElementById('chtx'), t=ta.value.trim(); if(!t) return;
    const b=document.getElementById('chgo'); b.disabled=true; b.textContent='보내는 중…';
    try{ const r=await fetch('/api/worship',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({action:'chat',date:NDATE,key:NKEY,text:t})});
      const j=await r.json(); if(j.ok){ ta.value=''; if(window.wsChatReload) wsChatReload(); } else alert('발송 실패: '+(j.error||'')); }
    catch(e){ alert('발송 실패 — 인터넷 연결을 확인해 주세요'); }
    b.disabled=false; b.textContent='챗발송'; };
  const hm=s=>{ try{ return new Date(s).toLocaleTimeString('ko-KR',{hour:'2-digit',minute:'2-digit',timeZone:'Asia/Seoul'}); }catch(e){ return ''; } };
  function draw(list){ const box=document.getElementById('chlist'); if(!box) return;
    const atEnd=box.scrollTop+box.clientHeight>=box.scrollHeight-4;
    box.replaceChildren();
    if(!list.length){ const n=document.createElement('div'); n.className='none'; n.textContent='확인할 메시지가 없습니다'; box.appendChild(n); return; }
    list.forEach(m=>{ const row=document.createElement('label'); row.className='cm';
      const cb=document.createElement('input'); cb.type='checkbox'; cb.title='확인 — 누르면 사라집니다';
      cb.onchange=async()=>{ row.style.opacity=.35; try{ await fetch('/api/worship',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({action:'chatdone',date:NDATE,key:NKEY,id:m.id})}); }catch(e){} if(window.wsChatReload) wsChatReload(); };
      const tx=document.createElement('div'); tx.textContent=m.text; const t=document.createElement('small'); t.textContent=hm(m.at); tx.appendChild(t);
      row.append(cb,tx); box.appendChild(row); });
    if(atEnd||box.dataset.n!==String(list.length)) box.scrollTop=box.scrollHeight; box.dataset.n=list.length; }
  addEventListener('wschat',e=>draw(e.detail||[]));
  addEventListener('DOMContentLoaded',()=>{ const ta=document.getElementById('chtx'), b=document.getElementById('chgo'); if(!ta) return;
    b.onclick=send;
    ta.addEventListener('keydown',e=>{ e.stopPropagation(); if(e.key==='Enter'&&!e.shiftKey&&!e.isComposing){ e.preventDefault(); send(); } });
    ta.addEventListener('keyup',e=>e.stopPropagation()); });
})();
"""


# ── ＋ 장 끼워 넣기 — 교회 소식·설교 뒤에 (2026-10-04 교장님) ──────────────────────────────
# 발표자 보기 아래 미리 보기 칸에 ＋ → ① 글 페이지(직접 쓰기) ② AI 디자인 페이지(문구를 쓰면 맥미니가 html 스킬로 글·도식 HTML,
# 배경은 템플릿 그대로) ③ 유튜브 페이지(프레젠테이션에서 꽉 차게 재생). 서버(ppt_pages)에 저장 — 프레젠테이션 창도 같이 바뀐다.
XP_CSS = """
#pvthumbs .th.add{display:flex;align-items:center;justify-content:center;aspect-ratio:16/9;border:3px dashed #475569;border-radius:8px;color:#cbd5e1;font:800 22px 'Pretendard Variable',sans-serif;cursor:pointer}
#pvthumbs .th.add:hover{border-color:#f6c76b;color:#f6c76b}
#pvthumbs .th .xpb{position:absolute;left:6px;top:6px;display:flex;gap:4px;z-index:3}#pvthumbs .th{position:relative}
#pvthumbs .th .xpb button{border:0;border-radius:6px;background:rgba(15,23,42,.85);color:#fff;font-size:14px;padding:3px 7px;cursor:pointer}
.gadd{display:none}body.grid:not(.pr) .gadd{display:flex;align-items:center;justify-content:center;aspect-ratio:16/9;border:3px dashed #475569;border-radius:8px;color:#cbd5e1;font:800 22px 'Pretendard Variable',sans-serif;cursor:pointer}
body.grid:not(.pr) .gadd:hover{border-color:#f6c76b;color:#f6c76b}
main .sl>.xpb{position:absolute;left:8px;top:8px;display:flex;gap:4px;z-index:5}body:not(.grid) main .sl>.xpb{display:none}
main .sl>.xpb button{border:0;border-radius:6px;background:rgba(15,23,42,.85);color:#fff;font-size:15px;padding:4px 8px;cursor:pointer}
#xpbox{position:fixed;inset:0;z-index:90;display:none;background:rgba(0,0,0,.6);align-items:center;justify-content:center}
#xpbox.on{display:flex}#xpbox .p{background:#fff;color:#111;border-radius:14px;width:min(820px,94vw);max-height:92vh;overflow:auto;padding:18px 20px;font:15px 'Pretendard Variable',sans-serif}
#xpbox h3{margin:0 0 10px}#xpbox .ch3{display:grid;grid-template-columns:repeat(3,1fr);gap:10px}
#xpbox .ch3 button{border:2px solid #e2e8f0;border-radius:12px;background:#f8fafc;padding:16px 10px;cursor:pointer;font:700 15px inherit;text-align:center;line-height:1.5}
#xpbox .ch3 button:hover{border-color:#f6c76b;background:#fffbeb}#xpbox .ch3 small{display:block;font-weight:400;color:#64748b;font-size:12.5px}
#xpbox label{display:block;font-weight:700;margin:10px 0 4px}#xpbox input,#xpbox textarea{width:100%;box-sizing:border-box;font:16px 'Pretendard Variable',sans-serif;border:1px solid #cbd5e1;border-radius:8px;padding:8px}
#xpbox textarea{min-height:180px;resize:vertical}#xpbox .hint{color:#64748b;font-size:13px;margin-top:6px}
#xpbox .b{display:flex;gap:8px;justify-content:flex-end;margin-top:14px}#xpbox .b button{font:700 15px inherit;border:0;border-radius:999px;padding:10px 18px;cursor:pointer}
#xpbox .ok{background:#1f2937;color:#fff}#xpbox .no{background:#e2e8f0}#xpbox .m{margin-right:auto;align-self:center;color:#b45309;font-size:13px}
#xpbox .ed{width:100%;aspect-ratio:16/9;border:1px solid #cbd5e1;border-radius:8px;background-size:cover;position:relative;overflow:hidden}
#xpbox .ed iframe{position:absolute;left:0;top:0;width:1920px;height:1080px;border:0;transform-origin:0 0}
"""

XP_JS = r"""
(function(){
const SECT={news:/^교회\s*소식$/, reading:/^성경\s*봉독$/, special:/^특\s*송$/, sermon:/^설교$/}, NAME={news:'교회 소식',reading:'성경 봉독',special:'특송',sermon:'설교'};   // 특송·성경 봉독 추가(2026-10-04)
// 그 주에만 끼운 순서(선교 보고 등, data ppt_extra)에도 늘 ＋ 장 넣기 (2026-10-04 교장님)
// XSECTS 는 이 스크립트보다 뒤에서 정해지므로 처음 쓸 때 읽는다 — 먼저 읽어 선교 보고에 ＋가 안 나왔던 일(2026-10-04)
let _S=null; const sects=()=>_S||(window.XSECTS?(window.XSECTS.forEach(t=>{ const k='x:'+t; SECT[k]=new RegExp('^'+t.replace(/[.*+?^${}()|[\]\\]/g,'\\$&').replace(/\s+/g,'\\s*')+'$'); NAME[k]=t; }),_S=Object.keys(SECT)):Object.keys(SECT));
let poll=null;
const anchor=sect=>[...document.querySelectorAll('#toc a')].find(a=>SECT[sect].test(a.textContent.trim()));
function startOf(sect){ const a=anchor(sect); return a?all.indexOf(document.getElementById(a.getAttribute('href').slice(1))):-1; }
function bg(){ const n=startOf('news'), nx=n>=0&&all[n+1]&&!all[n+1].classList.contains('xp')?D[+all[n+1].dataset.i]:null;
  const plain=d=>d&&d.img&&!d.html&&!d.pic&&d.texts.some(t=>t.c==='#ffffff'&&t.a==='l');      // 흰 글자 왼쪽 맞춤 = 갈색 내용 장(소식·봉독)
  if(plain(nx)) return nx.img; const any=D.find(plain); return any?any.img:(n>=0?D[+all[n].dataset.i].img:''); }
const ytId=u=>{ const m=String(u||'').match(/(?:youtu\.be\/|v=|embed\/|shorts\/|live\/)([\w-]{11})/); return m?m[1]:(/^[\w-]{11}$/.test(u)?u:''); };
// 📖 성경 봉독 장 = 기존 봉독 장 원칙 그대로 (2026-10-04 교장님): 한 장에 두 절, 절마다 위에 금빛 라벨(봉독대표·회중봉독 번갈아,
//   절 수가 홀수면 마지막 한 절은 「다함께 봉독」, 두 절 이하면 모두 「다함께 봉독」 — pptfill.read_label 과 같다). 긴 절은 글자만 줄인다.
function readBg(two){ const has=(d,w)=>d&&d.img&&!d.html&&d.texts.some(t=>t.t===w);
  const a=two?(D.find(d=>has(d,'봉독대표')&&has(d,'회중봉독'))):(D.find(d=>has(d,'다함께 봉독')&&!has(d,'회중봉독')));
  return a?a.img:(D.find(d=>has(d,'봉독대표')||has(d,'다함께 봉독'))||{}).img||bg(); }
function readLabel(k,n){ if(n<3) return '다함께 봉독'; if(k===n-1&&k%2===0) return '다함께 봉독'; return k%2===0?'봉독대표':'회중봉독'; }
function bibleD(d){ const V=(d.verses||[]).map(([v,t])=>v+' '+t), n=V.length, out=[];
  const fs=(x,room)=>{ const l=Math.ceil(x.length/24); return l<=room?55:Math.max(34,Math.floor(55*Math.sqrt(room/l))); };
  for(let k=0;k<n;k+=2){ const two=k+1<n, texts=[];
    texts.push({x:53,y:47,w:150,h:28,s:28,c:'#3b2a06',b:false,t:readLabel(k,n),a:'l'},
               {x:45,y:92,w:1350,h:0,s:fs(V[k],two?4:10),c:'#ffffff',b:true,t:V[k],a:'l',wrap:1,lh:1.17});
    if(two) texts.push({x:53,y:385,w:150,h:28,s:28,c:'#3b2a06',b:false,t:readLabel(k+1,n),a:'l'},
               {x:45,y:430,w:1350,h:0,s:fs(V[k+1],5),c:'#ffffff',b:true,t:V[k+1],a:'l',wrap:1,lh:1.17});
    out.push({img:readBg(two), bible:1, texts}); }
  return out; }
function pageD(p){ const d=p.data||{}, b=bg();
  if(p.type==='bible') return bibleD(d);
  if(p.type==='text'){ const tl=Math.max(1,Math.ceil((d.title||'').length/22));
    return {img:b, texts:[...(d.title?[{x:80,y:64,w:1280,h:0,s:62,c:'#ffffff',b:true,t:d.title,a:'l',wrap:true,lh:1.25}]:[]),
      {x:80,y:d.title?64+tl*78+36:70,w:1280,h:0,s:48,c:'#ffffff',b:false,t:d.body||'',a:'l',wrap:true,pre:true,lh:1.5}]}; }
  if(p.type==='yt') return {img:b, yt:d.id, pic:{url:'https://i.ytimg.com/vi/'+d.id+'/hqdefault.jpg',x:220,y:d.title?130:105,w:1000,h:562}, play:true,
      texts:d.title?[{x:220,y:40,w:1000,h:0,s:44,c:'#ffffff',b:true,t:d.title,a:'l',wrap:true,lh:1.2}]:[]};
  if(p.type==='html'&&p.status==='ready'&&d.html) return {img:b, html:d.html, texts:[]};
  return {img:b, texts:[{x:80,y:330,w:1280,h:0,s:50,c:'#f6c76b',b:true,a:'l',wrap:true,lh:1.35,
      t:p.status==='fail'?'⚠️ 만들지 못했습니다 — ✏️ 로 다시 시도해 주세요':'🛠 AI가 「'+(d.title||'새 페이지')+'」 페이지를 만드는 중… (1~2분)'}]};
}
window.xpPaint=(s,inn)=>{
  if(s.pic){ const im=document.createElement('img'); im.src=s.pic.url; im.alt='';
    im.style.cssText='position:absolute;left:'+s.pic.x+'px;top:'+s.pic.y+'px;width:'+s.pic.w+'px;height:'+s.pic.h+'px;object-fit:cover;border-radius:18px;box-shadow:0 10px 30px rgba(0,0,0,.4)'; inn.appendChild(im); }
  if(s.play){ const p=document.createElement('div'); p.className='xplay'; p.textContent='▶';
    p.style.cssText='position:absolute;left:'+(s.pic.x+s.pic.w/2-70)+'px;top:'+(s.pic.y+s.pic.h/2-70)+'px;width:140px;height:140px;border-radius:50%;background:rgba(220,38,38,.92);color:#fff;font-size:64px;display:flex;align-items:center;justify-content:center;padding-left:10px;box-sizing:border-box'; inn.appendChild(p); }
  if(s.html){ const f=document.createElement('iframe'); f.srcdoc=s.html;   // 스크립트는 맥미니가 빼고 저장(sandbox 를 걸면 글자가 안 그려졌다)
    f.style.cssText='position:absolute;left:0;top:0;width:1920px;height:1080px;border:0;transform:scale(.75);transform-origin:0 0;pointer-events:none;background:transparent'; inn.appendChild(f); }
};
// ── ＋ 메뉴 추가 (2026-10-08 교장님): 왼쪽 목차 BGM 위 ＋ → 제목을 넣으면 그 제목의 표지 한 장과 목차 줄이 생기고, 그 안에 ＋ 장 넣기
//   ppt_pages 에 type 'menu'(sect 'menu', data {title, after}) 로 저장, 그 안의 장은 sect 'm:<메뉴 id>'. 프레젠테이션 창도 같이 바뀐다.
const mainA=()=>TOC().filter(a=>!a.classList.contains('song')&&!a.classList.contains('xmenu'));
function coverD(title){ const L=mainA(), base=L.slice(1).concat(L.slice(0,1)).map(a=>D[+document.getElementById(a.getAttribute('href').slice(1)).dataset.i]).find(d=>d&&d.img&&!d.html&&(d.texts||[]).some(t=>t.s>=100));
  if(!base) return {img:bg(), texts:[{x:120,y:330,w:1200,h:0,s:110,c:'#ffffff',b:true,t:title,a:'c',wrap:true,lh:1.2}]};
  const big=base.texts.filter(t=>t.s>=100).sort((p,q)=>q.s-p.s)[0], cx=big.x+big.w/2, w=Math.min(1400,big.w*Math.max(1,title.length/Math.max(1,big.t.length)));
  return {img:base.img, cover:true, texts:[Object.assign({},big,{t:title,w:Math.round(w),x:Math.round(cx-w/2)})]}; }
function applyMenus(menus){
  document.querySelectorAll('#toc a.xmenu').forEach(a=>a.remove());
  sects(); if(!_S) _S=Object.keys(SECT); _S=_S.filter(k=>!/^m:/.test(k));
  for(const m of menus.slice().sort((p,q)=>p.id-q.id)){ const t=(m.data&&m.data.title||'').trim(); if(!t) continue;
    const k='m:'+m.id; SECT[k]=new RegExp('^'+t.replace(/[.*+?^${}()|[\]\\]/g,'\\$&')+'$'); NAME[k]=t; _S.push(k);
    const L=mainA(), af=m.data.after?L.findIndex(a=>a.textContent.trim()===m.data.after):-1;
    let nextA=af>=0?L[af+1]:L.find(a=>/^마침$/.test(a.textContent.trim().replace(/\s/g,'')));   // 기본: 마침 앞
    let at=nextA?all.indexOf(document.getElementById(nextA.getAttribute('href').slice(1))):-1; if(at<0) at=all.length;
    const di=D.length; D.push(coverD(t));
    const sec=document.createElement('section'); sec.className='sl xp'; sec.id='sxm'+m.id; sec.dataset.i=di; sec.dataset.xp=m.id; sec.dataset.sect=k;
    sec.innerHTML='<div class="in"></div><span class="no">+</span>';
    deck.insertBefore(sec, all[at]||document.getElementById('vbar')); all.splice(at,0,sec); io.observe(sec);
    const a=document.createElement('a'); a.href='#sxm'+m.id; a.className='xmenu'; a.textContent=t;
    a.onclick=ev=>{ ev.preventDefault(); a.blur(); if(window.bgmClose) bgmClose(); const el=document.getElementById('sxm'+m.id);
      if(document.body.classList.contains('pv')){ peek(all.indexOf(el),a); return; } if(document.body.classList.contains('grid')){ gridSec(a); return; }
      if(ONE()){ view(all.indexOf(el)); return; } el.scrollIntoView({behavior:'instant',block:'start'}); };
    const toc=document.getElementById('toc'); if(!toc) continue; const ref=nextA||toc.querySelector('a.xmadd')||toc.querySelector('a.bgm');
    if(ref&&ref.parentNode) ref.parentNode.insertBefore(a,ref); else toc.appendChild(a); }
}
function menuForm(pg){ const d=(pg&&pg.data)||{}, L=mainA();
  if(!box.isConnected) document.body.appendChild(box); box.classList.add('on');
  P().innerHTML='<h3>'+(pg?'✏️ 메뉴 고치기':'＋ 메뉴 추가')+'</h3><label>메뉴 제목</label><input class="t" maxlength="30" placeholder="예) 선교 보고 · 간증 · 광고" value="'+esc(d.title)+'">'
    +'<label>어디에 넣을까요</label><select class="w" style="width:100%;font:16px inherit;padding:8px;border:1px solid #cbd5e1;border-radius:8px"><option value="">맨 끝 (마침 앞)</option>'
    +L.filter(a=>!/^마침$/.test(a.textContent.trim().replace(/\s/g,''))).map(a=>{ const v=a.textContent.trim(); return '<option value="'+esc(v)+'"'+(v===d.after?' selected':'')+'>「'+esc(v)+'」 뒤</option>'; }).join('')+'</select>'
    +'<div class="hint">넣으면 이 제목의 표지 한 장이 생기고 왼쪽 목차에 메뉴가 붙습니다. 메뉴를 누른 뒤 모아 보기의 「＋ 장 넣기」로 글·AI 디자인·유튜브 장을 더할 수 있습니다.</div>'+foot(pg?'고치기':'넣기');
  const T=P().querySelector('.t'); setTimeout(()=>T.focus(),50); T.addEventListener('keydown',e=>{ if(e.key==='Enter'){ e.preventDefault(); P().querySelector('.ok').click(); } });
  wire(async()=>{ const t=T.value.trim(); if(!t) throw new Error('메뉴 제목을 넣어 주세요');
    if(mainA().some(a=>a.textContent.trim()===t&&!(pg&&a.getAttribute('href')==='#sxm'+pg.id))) throw new Error('같은 이름의 메뉴가 이미 있습니다');
    const j=await save({action:'page',id:pg&&pg.id,sect:'menu',type:'menu',data:{title:t,after:P().querySelector('.w').value}});
    if(!pg){ const a=[...document.querySelectorAll('#toc a.xmenu')].find(x=>x.textContent.trim()===t); if(a) setTimeout(()=>a.click(),50); } return j; }); }
window.xpMenuAdd=()=>menuForm(null);
// 목차에 「＋ 메뉴 추가」 — BGM 바로 위(발표자·편집 화면에서만, 프레젠테이션 창 X)
setTimeout(()=>{ const toc=document.getElementById('toc'); if(!toc||SCREEN||!NDATE||toc.querySelector('a.xmadd')) return;
  const a=document.createElement('a'); a.href='#'; a.className='xmadd'; a.textContent='＋ 메뉴 추가'; a.title='이 예배 PPT 에 순서(메뉴)를 더합니다';
  a.onclick=ev=>{ ev.preventDefault(); a.blur(); menuForm(null); };
  const b=toc.querySelector('a.bgm'); if(b) toc.insertBefore(a,b); else toc.appendChild(a); },0);
function applyPages(list){ const curEl=all[cur];
  all.filter(x=>x.classList.contains('xp')).forEach(x=>{ io.unobserve(x); x.remove(); all.splice(all.indexOf(x),1); });
  applyMenus(list.filter(p=>p.type==='menu'));
  for(const sect of sects()){ const s0=startOf(sect); if(s0<0) continue; let at=songRange(s0)[1];
    list.filter(p=>p.sect===sect).forEach(p=>[].concat(pageD(p)).forEach(dd=>{ const di=D.length; D.push(dd);   // 성경 봉독 한 건 = 여러 장
      const sec=document.createElement('section'); sec.className='sl xp'; sec.dataset.i=di; sec.dataset.xp=p.id; sec.dataset.sect=sect;
      sec.innerHTML='<div class="in"></div><span class="no">+</span>';
      deck.insertBefore(sec, all[at]||document.getElementById('vbar')); all.splice(at,0,sec); at++; io.observe(sec); })); }
  window.XP=list; const k=all.indexOf(curEl); if(k>=0) cur=k;
  fit(); thR=null;
  if(document.body.classList.contains('grid')&&window.wsGridAgain) wsGridAgain();   // 모아 보기 중이면 새 장·＋ 칸을 다시 그림
  if(document.body.classList.contains('pv')) go(cur,true); else if(SCREEN) show(cur,true); else if(ONE()) view(cur);
  clearTimeout(poll); if(list.some(p=>p.type==='html'&&(p.status==='wait'||p.status==='work'))) poll=setTimeout(loadPages,8000);
}
// 맥미니가 🎨 장을 다 만들거나 다시 만들면 프레젠테이션·발표자 화면이 30초 안에 저절로 새 것을 받는다 — 바뀐 것이 있을 때만 다시 그림 (2026-10-04 교장님)
let lastPages='';
async function loadPages(){ if(!NDATE) return; try{ const r=await fetch('/api/worship?pages='+NDATE,{cache:'no-store'}); if(!r.ok) return;
  const t=await r.text(); if(t===lastPages) return; lastPages=t; applyPages(JSON.parse(t)); }catch(e){} }
setInterval(loadPages,30000);
window.xpReload=loadPages;
async function save(body){ const r=await fetch('/api/worship',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(Object.assign({date:NDATE,key:NKEY},body))});
  const j=await r.json().catch(()=>({})); if(!j.ok) throw new Error(j.error||'저장 실패'); await loadPages(); if(bc) bc.postMessage({pages:1}); return j; }
// ── 미리 보기 칸: 교회 소식·설교 구간이면 ＋, 끼운 장에는 ✏️ 🗑
function xpBar(sect,sl){ const id=+sl.dataset.xp, pg=(window.XP||[]).find(p=>p.id===id); const bar=document.createElement('div'); bar.className='xpb';
  const isMenu=pg&&pg.type==='menu';
  const ed=document.createElement('button'); ed.textContent='✏️'; ed.title=isMenu?'메뉴 이름·자리 고치기':'고치기'; ed.onclick=ev=>{ ev.stopPropagation(); isMenu?menuForm(pg):open(sect,pg); };
  const del=document.createElement('button'); del.textContent='🗑'; del.title='빼기'; del.onclick=async ev=>{ ev.stopPropagation(); if(confirm(isMenu?'「'+pg.data.title+'」 메뉴와 그 안의 장을 모두 뺄까요?':'이 장을 뺄까요?')){ try{ await save({action:'pagedel',id}); if(isMenu&&document.body.classList.contains('grid')){ const f=TOC()[0]; if(f) gridSec(f); } }catch(err){ alert(err.message); } } };
  bar.append(ed,del); return bar; }
window.xpThumbs=(box,s,e)=>{
  const sect=sects().find(x=>startOf(x)===s); if(!sect||SCREEN) return;
  box.querySelectorAll('.th').forEach(t=>{ const sl=all[+t.dataset.k]; if(!sl||!sl.classList.contains('xp')) return; t.appendChild(xpBar(sect,sl)); });
  const add=document.createElement('div'); add.className='th add'; add.textContent='＋ 장 넣기'; add.title=NAME[sect]+' 뒤에 장 넣기'; add.onclick=()=>open(sect,null); box.appendChild(add);
};
// ── PPT 모아 보기에도 같은 ＋ (2026-10-07 교장님: 교회 소식·특송·설교 …) — 그 순서 장들 맨 끝에 ＋ 칸, 끼운 장에는 ✏️ 🗑
window.xpGrid=(s,e)=>{
  deck.querySelectorAll('.gadd,.sl>.xpb').forEach(x=>x.remove());
  const sect=sects().find(x=>startOf(x)===s); if(!sect||SCREEN||!NDATE) return;
  all.slice(s,e).filter(sl=>sl.classList.contains('xp')).forEach(sl=>sl.appendChild(xpBar(sect,sl)));
  const add=document.createElement('div'); add.className='gadd'; add.textContent='＋ 장 넣기'; add.title=NAME[sect]+' 뒤에 장 넣기';
  add.onclick=ev=>{ ev.stopPropagation(); open(sect,null); };
  deck.insertBefore(add, all[e]||document.getElementById('vbar'));
};
// ── 넣기·고치기 창
const box=document.createElement('div'); box.id='xpbox'; box.innerHTML='<div class="p"></div>';
box.addEventListener('keydown',e=>e.stopPropagation(),true); box.addEventListener('click',e=>{ if(e.target===box) box.classList.remove('on'); });
const P=()=>box.querySelector('.p');
const esc=t=>String(t||'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
function foot(okText){ return '<div class="b"><span class="m"></span><button class="no">닫기</button><button class="ok">'+okText+'</button></div>'; }
function wire(onOk){ P().querySelector('.no').onclick=()=>box.classList.remove('on');
  P().querySelector('.ok').onclick=async()=>{ const m=P().querySelector('.m'); m.textContent='저장 중…'; try{ await onOk(); box.classList.remove('on'); }catch(err){ m.textContent=err.message; } }; }
function open(sect,pg){ if(!box.isConnected) document.body.appendChild(box); box.classList.add('on');
  if(pg) return form(sect,pg.type,pg);
  if(sect==='reading'&&!open.more) return form(sect,'bible',null);   // 성경 봉독 ＋ = 곧바로 구절 찾기
  open.more=false;
  P().innerHTML='<h3>＋ '+NAME[sect]+' 뒤에 장 넣기</h3><div class="ch3">'
    +'<button data-t="text">📝 글 페이지<small>제목·내용을 직접 씁니다</small></button>'
    +'<button data-t="html">🎨 AI 디자인 페이지<small>문구를 쓰면 글·도식으로 꾸며 줍니다(배경은 템플릿)</small></button>'
    +'<button data-t="yt">▶ 유튜브 페이지<small>찾아서 고르면 프레젠테이션에서 꽉 차게 재생</small></button>'
    +(sect==='reading'?'<button data-t="bible">📖 성경 구절<small>찾으면 다함께 봉독 장으로 바로</small></button>':'')+'</div>'+foot('').replace('<button class="ok"></button>','');
  P().querySelector('.no').onclick=()=>box.classList.remove('on');
  P().querySelectorAll('.ch3 button').forEach(b=>b.onclick=()=>form(sect,b.dataset.t,null)); }
function form(sect,type,pg){ const d=(pg&&pg.data)||{};
  if(type==='bible'){   // 📖 구절 찾기 → 미리 보기 → 넣기 (AI 없이 즉시)
    P().innerHTML='<h3>📖 성경 봉독 — 구절 찾기</h3><label>구절 (예: 역대상 25:1-5 · 대상 25:1~5 · 시편 23편 · 요 3:16)</label>'
      +'<div class="yts"><input class="q" list="bbooks" autocomplete="off" placeholder="책 장:절-절" value="'+esc(d.ref)+'"><button type="button" class="go">🔎 찾기</button></div><datalist id="bbooks"></datalist>'
      +'<div class="bpv" style="max-height:340px;overflow-y:auto;margin:8px 0;font-size:14.5px;line-height:1.6;color:#111;background:#f8fafc;border-radius:8px;padding:8px 12px"></div>'
      +'<div class="hint">찾으면 바로 아래에 본문이 나옵니다. 「넣기」를 누르면 성경 봉독 뒤에 「다함께 봉독」 장으로 곧장 들어갑니다(한 장에 두 절 · 절마다 봉독대표·회중봉독 라벨, 홀수 마지막 절은 다함께 봉독).</div>'
      +foot(pg?'고치기':'넣기')+(pg?'':'<div style="text-align:right;margin-top:6px"><a href="#" class="oth" style="font-size:13px;color:#64748b">다른 장(글·AI·유튜브) 넣기 →</a></div>');
    const q=P().querySelector('.q'), V=P().querySelector('.bpv'); let got=null, tm=null;
    fetch('/api/worship?bible=books').then(r=>r.json()).then(b=>{ const dl=P().querySelector('#bbooks'); if(dl) dl.innerHTML=b.map(x=>'<option value="'+esc(x[0])+' ">'+esc(x[1])+' · '+x[2]+'장</option>').join(''); }).catch(()=>{});
    const find=async()=>{ const w=q.value.trim(); if(!/\d/.test(w)){ got=null; V.innerHTML=''; return; } V.textContent='찾는 중…';
      try{ const r=await fetch('/api/worship?bible='+encodeURIComponent(w)); const j=await r.json(); if(q.value.trim()!==w) return;
        if(!r.ok){ got=null; V.textContent=j.error||'찾지 못했습니다'; return; } got=j;
        const k=bibleD(j).length; V.innerHTML='<b>'+esc(j.ref)+'</b> · '+j.verses.length+'절 → '+k+'장<br>'+j.verses.map(([v,t])=>'<sup style="color:#b45309;font-weight:700">'+esc(v)+'</sup> '+esc(t)).join('<br>');
      }catch(e){ got=null; V.textContent='찾지 못했습니다 — 인터넷 연결을 확인해 주세요'; } };
    q.addEventListener('input',()=>{ clearTimeout(tm); tm=setTimeout(find,350); });
    q.addEventListener('keydown',e=>{ if(e.key==='Enter'){ e.preventDefault(); if(got&&got.ref&&q.value.trim()) P().querySelector('.ok').click(); else find(); } });
    P().querySelector('.go').onclick=find;
    const oth=P().querySelector('.oth'); if(oth) oth.onclick=e=>{ e.preventDefault(); open.more=true; open(sect,null); };
    if(!document.getElementById('bbcss')){ const st=document.createElement('style'); st.id='bbcss'; st.textContent='.yts{display:flex;gap:6px}.yts .q{flex:1}.yts .go{font:700 14px inherit;border:0;border-radius:8px;padding:0 14px;background:#1f2937;color:#fff;cursor:pointer}'; document.head.appendChild(st); }
    if(d.ref) find(); setTimeout(()=>q.focus(),50);
    wire(async()=>{ if(!got) await find(); if(!got) throw new Error('구절을 먼저 찾아 주세요 — 예) 역대상 25:1-5');
      return save({action:'page',id:pg&&pg.id,sect,type:'bible',data:{ref:got.ref,verses:got.verses}}); }); return; }
  if(type==='text'){ P().innerHTML='<h3>📝 글 페이지</h3><label>제목</label><input class="t" value="'+esc(d.title)+'"><label>내용</label><textarea class="c">'+esc(d.body)+'</textarea>'+foot(pg?'고치기':'넣기');
    wire(()=>save({action:'page',id:pg&&pg.id,sect,type:'text',data:{title:P().querySelector('.t').value.trim(),body:P().querySelector('.c').value}})); return; }
  if(type==='yt'){   // 찾기 → 결과에서 고르기 → 넣기 (2026-10-04 교장님)
    P().innerHTML='<h3>▶ 유튜브 페이지</h3><label>유튜브 찾기</label><div class="yts"><input class="q" placeholder="곡명·가수로 찾기 (예: 찬양 연속듣기)"><button type="button" class="go">🔎 찾기</button></div>'
      +'<div class="ytr"></div><label>고른 영상</label><div class="ysel">'+(d.id?'':'아직 고르지 않았습니다 — 위에서 찾아 누르거나 주소를 붙여 넣으세요')+'</div>'
      +'<input class="u" placeholder="또는 유튜브 주소 붙여 넣기 https://youtu.be/…" value="'+esc(d.id?'https://youtu.be/'+d.id:'')+'"><label>제목(선택)</label><input class="t" value="'+esc(d.title)+'">'
      +'<div class="hint">이 장이 프레젠테이션에 나오면 영상이 꽉 찬 화면으로 재생됩니다. 다른 장으로 넘기면 닫힙니다. (끊기지 않게 계속 틀 음악은 🎵 BGM 단추로)</div>'+foot(pg?'고치기':'넣기');
    if(!document.getElementById('ytscss')){ const st=document.createElement('style'); st.id='ytscss'; st.textContent=
      '.yts{display:flex;gap:6px}.yts .q{flex:1}.yts .go{font:700 14px inherit;border:0;border-radius:8px;padding:0 14px;background:#1f2937;color:#fff;cursor:pointer}'
      +'.ytr{display:grid;grid-template-columns:repeat(auto-fill,minmax(170px,1fr));gap:8px;max-height:300px;overflow-y:auto;margin:8px 0}'
      +'.ytr .r{display:flex;flex-direction:column;gap:3px;text-align:left;border:2px solid transparent;border-radius:8px;background:#f1f5f9;padding:4px;cursor:pointer;font:600 12.5px inherit;color:#111}'
      +'.ytr .r:hover{border-color:#94a3b8}.ytr .r.on{border-color:#dc2626;background:#fee2e2}.ytr .r img{width:100%;aspect-ratio:16/9;object-fit:cover;border-radius:6px}.ytr .r small{color:#64748b;font-weight:400}'
      +'.ysel{display:flex;gap:8px;align-items:center;font-size:13px;color:#475569;margin:2px 0 6px}.ysel img{width:120px;aspect-ratio:16/9;object-fit:cover;border-radius:6px}';
      document.head.appendChild(st); }
    const q=P().querySelector('.q'), R=P().querySelector('.ytr'), U=P().querySelector('.u'), T=P().querySelector('.t'), S=P().querySelector('.ysel');
    const pick=(id,title)=>{ U.value='https://youtu.be/'+id; if(title) T.value=title; S.innerHTML='<img src="https://i.ytimg.com/vi/'+esc(id)+'/mqdefault.jpg" alt=""><b>'+esc(title||id)+'</b>';
      R.querySelectorAll('.r').forEach(x=>x.classList.toggle('on',x.dataset.id===id)); };
    if(d.id) pick(d.id,d.title);
    const find=async()=>{ const w=q.value.trim(); if(!w) return; R.innerHTML='<div class="hint">찾는 중…</div>';
      try{ const list=await (await fetch('/api/worship?yts='+encodeURIComponent(w))).json();
        R.innerHTML=(list||[]).map(v=>'<button type="button" class="r" data-id="'+esc(v.id)+'" data-t="'+esc(v.title)+'"><img src="https://i.ytimg.com/vi/'+esc(v.id)+'/mqdefault.jpg" alt="" loading="lazy"><span>'+esc(v.title)+'</span><small>'+esc([v.ch,v.len,v.views&&('조회 '+v.views)].filter(Boolean).join(' · '))+'</small></button>').join('')||'<div class="hint">결과가 없습니다 — 다른 말로 찾아 보세요</div>';
      }catch(e){ R.innerHTML='<div class="hint">찾지 못했습니다 — 인터넷 연결을 확인하거나 주소를 붙여 넣어 주세요</div>'; } };
    P().querySelector('.go').onclick=find; q.addEventListener('keydown',e=>{ if(e.key==='Enter'){ e.preventDefault(); find(); } });
    R.onclick=e=>{ const b=e.target.closest('.r'); if(b) pick(b.dataset.id,b.dataset.t); };
    U.addEventListener('change',()=>{ const id=ytId(U.value.trim()); if(id) pick(id,T.value); });
    setTimeout(()=>q.focus(),50);
    wire(()=>{ const id=ytId(U.value.trim()); if(!id) throw new Error('영상을 찾아 고르거나 유튜브 주소를 넣어 주세요'); return save({action:'page',id:pg&&pg.id,sect,type:'yt',data:{id,title:T.value.trim()}}); }); return; }
  // html — 처음엔 문구, 다 만들어진 뒤엔 화면에서 글자를 바로 고친다
  if(pg&&d.html&&pg.status==='ready'){ P().innerHTML='<h3>🎨 AI 디자인 페이지 — 글자 고치기</h3><div class="ed" style="background-image:url('+(bg()?'slides/'+bg():'')+')"><iframe></iframe></div><div class="hint">글자를 눌러 바로 고칩니다. 내용을 크게 바꾸려면 「문구 바꿔 다시 만들기」.</div>'
      +'<div class="b"><span class="m"></span><button class="no">닫기</button><button class="re">문구 바꿔 다시 만들기</button><button class="ok">저장</button></div>';
    const ed=P().querySelector('.ed'), f=ed.querySelector('iframe'); f.srcdoc=d.html;
    const sc=()=>{ f.style.transform='scale('+(ed.clientWidth/1920)+')'; }; sc(); setTimeout(sc,50);
    f.onload=()=>{ try{ f.contentDocument.designMode='on'; }catch(e){} };
    P().querySelector('.re').onclick=()=>{ pg=Object.assign({},pg,{status:'wait'}); htmlForm(sect,pg,d); };
    wire(()=>{ const doc=f.contentDocument; doc.designMode='off'; return save({action:'page',id:pg.id,sect,type:'html',data:Object.assign({},d,{html:'<!doctype html>\n'+doc.documentElement.outerHTML})}); }); return; }
  htmlForm(sect,pg,d); }
function htmlForm(sect,pg,d){ P().innerHTML='<h3>🎨 AI 디자인 페이지</h3><label>제목</label><input class="t" value="'+esc(d.title)+'"><label>문구 (이 내용으로 글·도식 페이지를 만듭니다)</label><textarea class="c" placeholder="예) 추수감사절 홈커밍데이 — 10월 18일 주일 10:30, 사랑독채 펜션 별관. 예배 후 점심과 레크리에이션…">'+esc(d.body)+'</textarea><div class="hint">배경은 지금 PPT 템플릿 그대로 두고, 글과 간단한 도식(순서·일정·비교)으로 꾸밉니다. 맥미니가 1~2분 안에 만들어 이 자리에 넣습니다.</div>'+foot(pg?'다시 만들기':'AI로 만들기');
  wire(()=>{ const t=P().querySelector('.t').value.trim(), c=P().querySelector('.c').value.trim(); if(!c&&!t) throw new Error('문구를 써 주세요');
    return save({action:'page',id:pg&&pg.id,sect,type:'html',data:{title:t,body:c}}); }); }
// ── 프레젠테이션: 유튜브 장이면 꽉 차게 재생, 끼운 장 바뀌면 다시 불러오기
let ytOn=false;
// ＋ 로 넣은 유튜브 장은 BGM 창이 아니라 프레젠테이션 자체에서 꽉 차게 — 다른 장으로 넘기면 닫힌다 (2026-10-04 교장님)
// 덮개(bgmFS)가 아니라 그 장 안에 영상을 끼운다 — 장을 넘기면 영상도 같이 빠지고 다음 장이 나온다 (2026-10-04 교장님)
let ytEl=null;
function ytCheck(){ const sl=all[cur], d=sl&&D[+sl.dataset.i];
  if(ytEl&&(!d||!d.yt||ytEl.parentNode!==sl.querySelector('.in'))){ ytEl.remove(); ytEl=null; ytOn=false; }
  if(d&&d.yt&&document.body.classList.contains('pr')&&!ytEl){ const inn=sl.querySelector('.in'); paint(sl);
    const f=document.createElement('iframe'); f.className='ytin';
    f.src='https://www.youtube-nocookie.com/embed/'+d.yt+'?autoplay=1&rel=0&modestbranding=1&playsinline=1';
    f.allow='autoplay; encrypted-media; fullscreen'; f.referrerPolicy='strict-origin-when-cross-origin';
    f.style.cssText='position:absolute;left:0;top:0;width:1440px;height:810px;border:0;z-index:6;background:#000';
    inn.appendChild(f); ytEl=f; ytOn=true; } }
addEventListener('DOMContentLoaded',()=>{
  const _s=show; show=function(i,q){ _s(i,q); ytCheck(); };
  if(bc){ const prev=bc.onmessage; bc.onmessage=e=>{ const m=e.data||{}; if(m.pages){ loadPages(); return; } if(prev) prev(e); }; }
  deck.addEventListener('click',e=>{ const sl=e.target.closest('.sl.xp'); if(!sl||document.body.classList.contains('pr')) return; const d=D[+sl.dataset.i]; if(!d||!d.yt) return;
    const inn=sl.querySelector('.in'); if(inn.querySelector('iframe')) return; const f=document.createElement('iframe');   // 이 화면 안에서 재생(다른 장으로 가면 닫힘)
    f.src='https://www.youtube-nocookie.com/embed/'+d.yt+'?autoplay=1&rel=0'; f.allow='autoplay; encrypted-media; fullscreen'; f.allowFullscreen=true;
    f.style.cssText='position:absolute;left:'+d.pic.x+'px;top:'+d.pic.y+'px;width:'+d.pic.w+'px;height:'+d.pic.h+'px;border:0;border-radius:18px;z-index:5'; inn.appendChild(f); });
  loadPages(); });
})();
"""


def _news_globals(date: str | None) -> str:
    """교회 소식 고치기 저장 열쇠(그 주 날짜로) — api/worship.js 의 news 와 같은 계산."""
    import hmac, hashlib
    sec = next((ln.split("=", 1)[1].strip() for ln in (Path.home() / "dev/daily-briefing/.env").read_text().splitlines() if ln.startswith("JUBO_SECRET=")), "")
    key = hmac.new(sec.encode(), f"news:{date}".encode(), hashlib.sha256).hexdigest()[:32] if sec and date else ""
    try:
        xs = [x.get("title", "") for x in json.loads((HERE / "data" / f"{date}.json").read_text()).get("ppt_extra", []) if x.get("title")]
    except Exception:
        xs = []
    return f"const NDATE='{date or ''}',NKEY='{key}';window.XSECTS={json.dumps(xs, ensure_ascii=False)};\n"


def render(slides: list[dict], title: str, dl: str, songs: list[str] | None = None, date: str | None = None) -> str:
    data = [{"img": s["img"], "texts": s["texts"]} for s in slides]
    sl = "".join(f'<section class="sl" id="s{s["n"]}" data-i="{k}"><div class="in"></div><span class="no">{s["n"]}</span></section>'
                 for k, s in enumerate(slides))
    toc = "".join(f'<a href="#s{n}" class="{"song" if name.startswith("♪") else ""}">{html.escape(name)}</a>' for n, name in outline(slides, list(songs or [])))
    import re as _re, wsnav                      # 공통 상단 메뉴 (2026-10-03 교장님: 쪽마다 상단 통일)
    m = _re.search(r"\d+월 \d+일", title)
    dls = [(("📊 PPT 받기" if h.endswith(".pptx") else "📕 PDF 받기"), h) for h in _re.findall(r'href="([^"]+)"', dl or "")]
    note = "" if dls else _re.sub(r"<[^>]+>", "", dl or "")
    sub = ('<button class="sub on" id="b-grid" onclick="if(document.body.classList.contains(\'pv\'))endPV();wsGrid()">📑 PPT</button>'   # 2026-10-07 교장님: PPT / 예배용(두 화면) / BGM
           '<button class="sub" id="b-pv" onclick="present()">▶ 예배용(두 화면)</button>')
    tail = f'<span class="note">{html.escape(note)}</span>' if note.strip() else ""
    if not date and m:                            # 날짜를 안 받았으면 제목(「10월 11일」)에서
        import datetime as dt
        mo, da = map(int, re.findall(r"\d+", m.group(0))); date = f"{dt.date.today().year}-{mo:02d}-{da:02d}"
    items = bgm(date) if date else {}
    import svc as _svc
    first_bgm = "기도용 CCM(경건)" if (date and _svc.service_of(date)["id"] == "dawn") else "경쾌한 CCM"   # 새벽예배는 경건한 기도 찬양부터
    if items:
        toc += '<a href="#bgm" class="bgm">🎵 BGM</a>'
        sub += '<button class="sub" onclick="bgmOpen()">🎵 BGM</button>'
    sub += tail
    toc += CHAT_HTML                              # 💬 예배팀 소통(목차 아래) — 2026-10-04 교장님
    # 「예배용 PPT 열기」(모음 쪽)는 ppt.html#pv — 바로 발표자 보기로 열고, 📺 프레젠테이션 열기 한 번으로 두 번째 모니터 (창은 눌러야 열린다)
    nav_pv = ("<script>addEventListener('DOMContentLoaded',()=>{ if(D.some(s=>s.texts.some(t=>t.css&&/Serif|Cormorant/.test(t.css)))){const l=document.createElement('link');l.rel='stylesheet';"
              "l.href='https://fonts.googleapis.com/css2?family=Noto+Serif+KR:wght@400;600;900&family=Cormorant+Garamond:ital,wght@0,500;1,500&display=swap';document.head.appendChild(l);"
              "(document.fonts?document.fonts.ready:Promise.resolve()).then(()=>{document.querySelectorAll('.sl').forEach(x=>{delete x.dataset.p;x.querySelectorAll('.tx').forEach(e=>e.remove());});document.querySelectorAll('.sl.cur').forEach(x=>paint(x));});} });</script>")
    nav_pv += "<script>addEventListener('load',()=>{ if(location.hash==='#pv'&&!new URLSearchParams(location.search).has('screen')) enterPV(0); });</script>"
    # 2단 밑줄(켜짐)은 지금 보기를 따라간다 — 발표자 보기면 ▶ 예배용, 아니면 한 장씩·목록 (2026-10-04 교장님)
    nav_pv += ("<script>addEventListener('DOMContentLoaded',()=>{ const sync=()=>{ const c=document.body.classList, pv=c.contains('pv');"
               "const set=(id,on)=>{ const b=document.getElementById(id); if(b) b.classList.toggle('on',on); };"
               "set('b-pv',pv); set('b-grid',!pv); };"
               "new MutationObserver(sync).observe(document.body,{attributes:true,attributeFilter:['class']}); sync(); });</script>")
    nav = nav_pv + wsnav.nav("ppt", m.group(0) if m else title, sub, dls, day=date) + f"<script>{wsnav.JS}</script>" + bgm_html(items, first_bgm) + f"<script>{CHAT_JS}</script><script>{XP_JS}</script>"
    return PAGE.format(title=html.escape(title), n=len(slides), slides=sl, toc=toc, dl=dl, nav=nav, navcss=wsnav.CSS + BGM_CSS + CHAT_CSS + XP_CSS,
                       data=json.dumps(data, ensure_ascii=False, separators=(",", ":")), pvjs=_news_globals(date) + PVJS)


def rewrap(slides: list[dict], pptx: Path | None) -> list[dict]:
    """PDF 는 줄마다 끊겨 있다 — 원본 PPT 의 한 문단이 여러 줄로 꺾인 것은 다시 한 덩어리로 이어,
    화면 글꼴 폭에 맞춰 자연스럽게 줄이 바뀌게 한다(2026-10-03 교장님: 오른쪽 여백이 남지 않게, 교회소식·봉독 본문).
    왼쪽 맞춤 글만. 이어 붙인 줄들이 원본 문단과 글자 그대로 같을 때만 합친다."""
    if not pptx or not Path(pptx).exists(): return slides
    from pptx import Presentation
    prs = Presentation(str(pptx))
    src = [x for x in prs.slides if not any(sh.name.startswith("songppt|") for sh in x.shapes)]   # PDF 는 곡 장을 뺀 사본에서 나온다
    flat = lambda t: re.sub(r"\s", "", t)
    for k, sl in enumerate(slides):
        if k >= len(src): break
        paras = [flat("".join(r.text for r in pa.runs)) for sh in src[k].shapes if sh.has_text_frame
                 for pa in sh.text_frame.paragraphs]
        paras = [x for x in paras if x]
        out, i, T = [], 0, sl["texts"]
        while i < len(T):
            t = dict(T[i]); j = i
            if t.get("a") == "l":
                acc = flat(t["t"])
                while (j + 1 < len(T) and T[j + 1].get("a") == "l" and abs(T[j + 1]["x"] - t["x"]) < 6 and abs(T[j + 1]["s"] - t["s"]) < 1
                       and acc not in paras and any(p.startswith(acc + flat(T[j + 1]["t"])) for p in paras)):
                    j += 1; acc += flat(T[j]["t"])
                if j > i and acc in paras:
                    seg = T[i:j + 1]
                    t["t"] = " ".join(x["t"].strip() for x in seg)
                    t["b"] = any(x["b"] for x in seg)
                    t["lh"] = round((seg[-1]["y"] - seg[0]["y"]) / (len(seg) - 1) / t["s"], 2)
                    t["w"] = max(max(x["w"] for x in seg), sl["w"] - 2 * t["x"])     # 칸 폭(좌우 같은 여백)까지
                    t["h"] = seg[-1]["y"] + seg[-1]["h"] - t["y"]
                    t["wrap"] = 1
                else:
                    j = i
            out.append(t); i = j + 1
        sl["texts"] = out
        sl["plain"] = " ".join(x["t"] for x in out)
    return slides


_flat40 = lambda s: re.sub(r"\s", "", s["plain"])[:40]


def cover_kind(s: dict) -> str | None:
    """찬양 표지 장이면 「경배」·「결단」(가사 장처럼 글이 긴 장은 아니다)."""
    f = _flat40(s)
    if ("찬양과경배" in f or "찬양과결단" in f) and len(re.sub(r"\s", "", s["plain"])) < 160:
        return "결단" if "찬양과결단" in f else "경배"
    return None


def collapse_covers(slides: list[dict]) -> list[dict]:
    """같은 찬양 표지가 연달아 있으면 1장만 남긴다 — 템플릿에 곡 자리마다 표지가 있던 것
    (2026-10-07 교장님: 찬양과경배 표지는 1개, 곡은 그 아래 여러 개). 사이에 가사 장이 있으면 그대로 둔다."""
    res = []
    for s in slides:
        k = cover_kind(s)
        if k and res and cover_kind(res[-1]) == k:
            continue
        res.append(s)
    for n, s in enumerate(res, 1): s["n"] = n
    return res


def song_plan(kinds: list[str], groups: dict) -> list[list]:
    """표지마다 넣을 곡. 경배 표지가 둘이면 도입곡 → 첫째(사도신경 앞), 진행곡 → 둘째. 하나면 둘 다 그 아래.
    적용송 → 결단 표지(없으면 마지막 경배 표지). 표지가 결단뿐이면 모두 그 아래."""
    plan = [[] for _ in kinds]
    W = [i for i, k in enumerate(kinds) if k == "경배"]; K = [i for i, k in enumerate(kinds) if k == "결단"]
    intro, main, apply = list(groups.get("intro", [])), list(groups.get("main", [])), list(groups.get("apply", []))
    if len(W) >= 2: plan[W[0]] += intro; plan[W[1]] += main
    elif W: plan[W[0]] += intro + main
    elif K: plan[K[0]] += intro + main
    if K: plan[K[0]] += apply
    elif W: plan[W[-1]] += apply
    return plan


def add_song_frames(slides: list[dict], date: str, out: Path | None) -> list[dict]:
    """찬양과경배·찬양과결단 표지 뒤에 그 순서의 곡을 모두 끼운다 — 곡마다 곡별 PPT 장, 없으면 틀 한 장(곡 제목 + 악보)
    (2026-10-03 교장님: 곡마다 자리가 있게 / 2026-10-07: 표지는 순서마다 1장, 곡은 그 아래 여러 개).
    표지 뒤에 이미 가사 장이 있으면 곡은 손대지 않는다."""
    slides = collapse_covers(slides)
    try:
        d = json.loads((HERE / "data" / f"{date}.json").read_text()).get("songs", {})
    except Exception:
        return slides
    lab = lambda g, i: {"intro": "도입곡", "apply": "적용송"}.get(g) or str(i + 1)
    groups = {g: [(lab(g, i), x) for i, x in enumerate(d.get(g, []))] for g in ("intro", "main", "apply")}
    # 다음 장이 곧바로 다른 순서 표지 = 가사 장 없음. 예배순서에서 온 표지(합심기도·간증 등 _sec·cover)도 순서 표지(2026-10-08)
    is_sec = lambda s: any(x in _flat40(s) for x in SECTIONS) or bool(s.get("_sec")) or bool(s.get("cover"))
    divs = [k for k, s in enumerate(slides) if cover_kind(s)]
    empty = [k for k in divs if k + 1 >= len(slides) or is_sec(slides[k + 1])]
    if not any(groups.values()) or not divs or len(empty) < len(divs) or out is None:   # 가사 장이 이미 들어 있는 PPT
        return slides
    from PIL import Image
    plan = dict(zip(divs, song_plan([cover_kind(slides[k]) for k in divs], groups)))
    W, H = int(slides[0]["w"]), int(slides[0]["h"])
    for f in out.glob("song_*.jpg"): f.unlink()
    res, si = [], 0
    for k, s in enumerate(slides):
        res.append(s)
        for lab_, x in plan.get(k, []):
            si += 1
            name_ = x.get("title") or "곡 미정"
            got = song_slides(x.get("title", ""), si, W, H, out)       # 곡별 PPT DB 에 있으면 그 곡의 장을 모두
            if got:
                got[0]["start"] = name_; res += got; continue
            name = f"frame_{si:02d}.jpg"
            cv = Image.new("RGB", (W, H), "white")
            if x.get("img") and (HERE / x["img"]).exists():
                im = Image.open(HERE / x["img"]).convert("RGB")
                bw, bh = W - 80, H - 130
                r = min(bw / im.width, bh / im.height); im = im.resize((max(1, int(im.width * r)), max(1, int(im.height * r))))
                cv.paste(im, ((W - im.width) // 2, 110 + (bh - im.height) // 2))
            cv.save(out / name, "JPEG", quality=85)
            title = ("♬ " + (lab_ + ". " if lab_.isdigit() else lab_ + " · ") + name_)
            texts = [{"x": 40, "y": 30, "w": W - 80, "h": 60, "s": 48, "c": "#6b5444", "b": True, "t": title, "a": "c"}]
            if not x.get("img"):
                texts.append({"x": 40, "y": H // 2, "w": W - 80, "h": 40, "s": 32, "c": "#9ca3af", "b": False, "t": "악보·가사 준비 중", "a": "c"})
            res.append({"n": 0, "w": W, "h": H, "img": name, "texts": texts, "plain": x.get("title", ""), "hidden": x.get("title", ""), "start": name_})
    for n, s in enumerate(res, 1): s["n"] = n
    return res


def song_slides(title: str, si: int, W: int, H: int, out: Path) -> list[dict]:
    """곡별 PPT DB(songppt/db)에서 곡을 찾아 예배 PPT 장으로 — 악보 그림 + 자막 바탕(검정)은 그림에, 자막 글은 글자로
    (2026-10-03 교장님 지시: 콘티 순서대로 찬양 자리에 곡 PPT 를 자동으로)."""
    import base64, io
    from PIL import Image, ImageDraw
    sys.path.insert(0, str(HERE)); sys.path.insert(0, str(HERE.parent))
    try:
        import songppt, store
        t = store.song_find(title)                                # 곡별 PPT 는 예배 DB + 그림 창고에서
        song = store.song_load(t) if t else None
        if not song: return []
    except Exception:
        return []
    CW, CH = 1920, 1080; kx, ky = W / CW, H / CH
    res = []
    for j, sl in enumerate(song["slides"], 1):
        cv = Image.new("RGB", (CW, CH), "white")
        im = Image.open(io.BytesIO(base64.b64decode(sl["img"]))).convert("RGBA")
        x, y, w, h = sl["pic"]
        im = im.resize((max(1, w), max(1, h)))
        cv.paste(im, (x, y), im)
        texts = []
        lay = songppt.sub_layout(sl.get("sub"))           # 자막 통일: 아래 띠 20% — 위 절반 러시아어(노랑)·아래 절반 영어(흰색)
        if lay:
            bx, by, bw, bh = lay["box"]
            ImageDraw.Draw(cv).rectangle([bx, by, bx + bw, by + bh], fill=lay["fill"])
            for ln in lay["lines"]:
                texts.append({"x": round(ln["x"] * kx), "y": round(ln["y"] * ky), "w": round(ln["w"] * kx), "h": round(ln["size"] * ky),
                              "s": round(ln["size"] * ky, 1), "c": ln["c"], "b": True, "t": ln["t"], "a": "c"})
        name = f"song_{si:02d}_{j:02d}.jpg"
        cv.save(out / name, "JPEG", quality=85)
        res.append({"n": 0, "w": W, "h": H, "img": name, "texts": texts, "plain": t, "hidden": t, "song": t})
    return res


def add_extra(slides: list[dict], date: str, out: Path | None = None) -> list[dict]:
    """그 주에만 끼우는 순서 장 — data/<날짜>.json 의 ppt_extra (2026-10-03 교장님 지시: 특송 다음 「선교 보고 · 디마 선교사」).
    [{"after": "SpecialPraise", "base": "ScriptureReading", "label": "Mission Report", "title": "선교 보고", "name": "디마 선교사"}]
    base 장(제목 상자 크기가 맞는 순서 표지)의 배경을 빌려 글자만 바꾼다. after 장 바로 뒤에 넣는다."""
    try:
        extra = json.loads((HERE / "data" / f"{date}.json").read_text()).get("ppt_extra", [])
    except Exception:
        return slides
    return insert_extra(slides, extra, out, date)


def make_cover(base: dict, title: str = "", label: str = "", name: str = "", sub: str = "") -> dict:
    """순서 표지 한 장 — base(같은 디자인 표지)의 배경을 빌려 글자만 바꾼다(그 주에만 끼우는 순서·예배순서에서 새로 넣은 순서)."""
    s = json.loads(json.dumps(base)); s.pop("_sec", None); texts = []
    for t in s["texts"]:
        if t["s"] >= 100: val = title
        elif t["s"] < 30 and t["y"] < 300: val = label
        elif 40 <= t["s"] < 100: val = name
        else: val = sub
        if not val: continue
        cx = t["x"] + t["w"] / 2; w = t["w"] * max(1.0, len(val) / max(1, len(t["t"])))
        texts.append({**t, "t": val, "w": int(w), "x": int(cx - w / 2)})
    s.update(texts=texts, plain=" ".join(t["t"] for t in texts), hidden="", cover=True)
    return s


def insert_extra(slides: list[dict], extra: list[dict], out: Path | None, date: str = "") -> list[dict]:
    flat = lambda s: re.sub(r"\s", "", s["plain"])
    big = lambda s: next((re.sub(r"\s", "", t["t"]) for t in s.get("texts", []) if t.get("s", 0) >= 100), "")
    for x in extra:
        same = next((i for i, s in enumerate(slides) if big(s) and big(s) == re.sub(r"\s", "", x.get("title", ""))), None)
        if same is not None:           # 예배순서에 이미 이 순서가 있다(PPT → 예배순서로 적힌 뒤) — 표지는 두지 않고 보고 자료 장만 그 표지 뒤에 (2026-10-07)
            k, add = same, []
        else:
            base = next((s for s in slides if flat(s).startswith(x.get("base", "ScriptureReading"))), None)
            k = next((i for i, s in enumerate(slides) if flat(s).startswith(x["after"])), None)
            if base is None or k is None: continue
            s = make_cover(base, x.get("title", ""), x.get("label", ""), x.get("name", ""), x.get("sub", ""))
            s["_sec"] = {"t": x.get("title", ""), **({"who": x["name"]} if x.get("name") else {})}   # 예배순서에 다시 적힐 줄
            add = [s]
        # 보고 자료(PDF)가 있으면 장마다 그림 한 장씩 표지 뒤에 이어 붙인다 — 글자까지 원본 그대로 (2026-10-03 디마 선교사 몽골 단기선교 보고)
        src = HERE / "extra" / date / x["pdf"] if x.get("pdf") and date else None
        if src and src.exists() and out is not None:
            import fitz
            doc = fitz.open(src); tag = re.sub(r"\W", "", x.get("name", "extra"))[:8] or "extra"
            for i, pg in enumerate(doc, 1):
                name = f"x_{tag}_{i:02d}.jpg"
                pg.get_pixmap(matrix=fitz.Matrix(1920 / pg.rect.width, 1920 / pg.rect.width)).save(str(out / name), jpg_quality=88)
                add.append({"n": 0, "w": slides[k]["w"], "h": slides[k]["h"], "img": name, "texts": [], "plain": "", "hidden": ""})
        slides[k + 1:k + 1] = add
    for n, s in enumerate(slides, 1): s["n"] = n
    return slides


def song_titles(date: str) -> list[str]:
    """그 주 악보집 곡 순서(도입곡 → 1~N → 적용송). 없으면 악보 보관함의 「쓴 날」 기록에서."""
    try:
        d = json.loads((HERE / "data" / f"{date}.json").read_text()).get("songs", {})
        t = [x.get("title", "") for g in ("intro", "main", "apply") for x in d.get(g, [])]
        if any(t): return t
        bank = json.loads((HERE / "score_bank.json").read_text())["songs"]
        order = {"도입곡": 0, "적용송": 99}
        hits = []
        for title, rec in bank.items():
            for sc in rec["scores"].values():
                for u in sc.get("used", []):
                    if u.startswith(date + " "):   # 같은 날 다른 예배(날짜-열쇠)와 섞이지 않게
                        lab = u.split(" ", 1)[1]
                        hits.append((order.get(lab, int(re.sub(r"\D", "", lab) or 50)), title))
        return [t for _, t in sorted(hits)]
    except Exception:
        return []


def make(pdf: Path, date: str, sid: str, pptx: Path | None = None) -> Path:
    folder = SITE / "d" / sid; folder.mkdir(parents=True, exist_ok=True)
    slides = parse(pdf, folder / "slides")
    slides = rewrap(slides, pptx)
    import fitz                                   # 모음 쪽(jegok_worship) 「오늘」 칸에 쓸 첫 장 그림(글자까지 그대로)
    fitz.open(pdf)[0].get_pixmap(dpi=110).save(str(folder / "cover.jpg"))
    slides = add_song_frames(slides, date, folder / "slides")
    slides = add_extra(slides, date, folder / "slides")
    shutil.copy(pdf, folder / "worship.pdf")
    dl = '<a href="worship.pdf" download="{0} 주일예배 PPT.pdf">⬇ PDF</a>'.format(date)
    if pptx:
        shutil.copy(pptx, folder / "worship.pptx")
        dl = f'<a href="worship.pptx" download="{date} 주일예배 PPT.pptx">⬇ PPT</a>' + dl
    title = f"{int(date[5:7])}월 {int(date[8:10])}일 주일예배 PPT"
    (folder / "ppt.html").write_text(render(slides, title, dl, song_titles(date), date=date))
    return folder


def fetch(date: str) -> str:
    """그 주 드라이브 폴더의 「…주일예배 PPT」(.pptx)를 받아 worship_ppt/<날짜>.pptx·.pdf 로 둔다 (2026-10-03 교장님 지시).
    PDF 는 구글 슬라이드로 잠깐 바꿔 내보낸 뒤 그 임시본을 지운다. 드라이브 것이 더 새것일 때만 다시 받는다."""
    import datetime as dt, os
    sys.path.insert(0, str(HERE.parent)); sys.path.insert(0, str(HERE))
    import prep, weekly
    g = prep.G(prep.access_token()); fo = weekly.week_folder(g, dt.date.fromisoformat(date))
    if not fo: return "드라이브 폴더 없음"
    r = g.get(f"{prep.DRIVE}/files", q=f"'{fo['id']}' in parents and trashed=false and name contains '주일예배 PPT'",
              fields="files(id,name,mimeType,modifiedTime)", includeItemsFromAllDrives="true")
    f = next((x for x in sorted(r.get("files", []), key=lambda x: x["modifiedTime"], reverse=True)
              if x["mimeType"].endswith("presentationml.presentation")), None)
    if not f: return "주일예배 PPT 없음"
    px, pdf = HERE / "worship_ppt" / f"{date}.pptx", HERE / "worship_ppt" / f"{date}.pdf"
    mt = dt.datetime.fromisoformat(f["modifiedTime"].replace("Z", "+00:00")).timestamp()
    if pdf.exists() and px.exists() and px.stat().st_mtime >= mt: return "주일예배 PPT 그대로"
    px.parent.mkdir(exist_ok=True); px.write_bytes(g.download(f["id"]))
    import pptfill                                   # 곡 장(악보)까지 든 PPT 는 구글 내보내기 한도를 넘는다 — 곡 장을 뺀 사본으로 PDF
    base = pptfill.base_copy(px, px.with_name(px.stem + ".base.pptx"))
    try:
        pptfill.to_pdf(base, pdf)
    finally:
        base.unlink(missing_ok=True)
    os.utime(px, (mt, mt))
    return f"주일예배 PPT 받음 · {f['name']}"


if __name__ == "__main__":
    sys.path.insert(0, str(HERE))
    if sys.argv[1] == "fetch":
        print(fetch(sys.argv[2])); sys.exit(0)
    import publish
    a = sys.argv[1:]
    pptx = Path(a[a.index("--pptx") + 1]) if "--pptx" in a else None
    print(make(Path(a[0]), a[1], publish.sid_for(a[1]), pptx))
