#!/usr/bin/env python3
"""곡별 PPT — 주일예배 PPT 에서 곡을 잘라 곡마다 「곡명.html」 로 (2026-10-03 교장님 승인).

예배 PPT 의 곡 슬라이드는 이미 층으로 나뉘어 있다:
  · 악보 층  = PNG 그림 한 장(악보 + 한글 가사, 화면 위 0~91%)
  · 자막 층  = 아래 검은 바탕 글상자(러시아어 / 영어 줄) — 글자로 뽑아 편집할 수 있게 둔다
  · 절 단추  = 맨 아래 「1.사람을 보며…」 단추(누르면 그 절로)
곡 경계는 악보 그림 왼쪽 위 곡 제목 부분이 바뀌는 곳, 이름은 그 주 콘티(악보집) 곡 순서.

  python3 accomp/songppt.py extract <예배PPT.pptx> 2026-09-27 [--only 곡명] [--upload]
      → db/<곡명>/song.json + 그림, out/<곡명>.html, (--upload) 드라이브 「곡별 PPT」 폴더
  python3 accomp/songppt.py save < 수정.json     자막 고친 것 저장(docsave /songppt/save 가 부른다)
"""
from __future__ import annotations

import base64, html, io, json, re, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE / "songppt"
DB, OUT = ROOT / "db", ROOT / "out"
DRIVE_ROOT = "1feStbHToyPB8MHrPlRvPuVZkf_UCZVt2"     # 00 주일예배 콘티
DRIVE_FOLDER_NAME = "곡별 PPT"
SW, SH = 1920, 1080                                   # 화면 좌표(16:9)


# ── 판정 규칙 ───────────────────────────────────────────────────
def different(a: str, b: str) -> bool:
    if len(a) == len(b) == 256:                       # 제목 부분 그림 지문(평균 해시)
        return sum(x != y for x, y in zip(a, b)) > 40
    return a != b


def segments(slides: list[dict]) -> list[tuple[int, int]]:
    """곡 장이 이어지는 구간. 곡이 아닌 장(표지·사도신경)에서 끊고, 제목 그림이 바뀌면 새 곡.
    붙어 있는 한 장짜리 구간(곡 첫 장 그림이 다른 경우)은 바로 뒤 구간에 합친다."""
    segs: list[list[int]] = []
    prev = None
    for s in slides:
        if not s["song"] or s.get("creed"):
            prev = None; continue
        if prev is None or prev["n"] != s["n"] - 1 or different(prev["title_sig"], s["title_sig"]):
            segs.append([s["n"], s["n"]])
        else:
            segs[-1][1] = s["n"]
        prev = s
    out: list[list[int]] = []
    i = 0
    while i < len(segs):
        a = segs[i]
        if a[0] == a[1] and i + 1 < len(segs) and segs[i + 1][0] == a[1] + 1:
            out.append([a[0], segs[i + 1][1]]); i += 2; continue
        out.append(a); i += 1
    return [tuple(x) for x in out]


def name_segments(segs: list[tuple[int, int]], songs: list[str]) -> list[tuple[str, int, int]]:
    if len(segs) == len(songs):
        return [(t, a, b) for t, (a, b) in zip(songs, segs)]
    return [(f"확인필요-{a}", a, b) for a, b in segs]          # 개수가 안 맞으면 이름을 지어내지 않는다


def file_name(title: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r'[\\/:*?"<>|]', " ", title)).strip() + ".html"


# ── PPT 읽기 ────────────────────────────────────────────────────
def _hex(font_or_fill) -> str | None:
    try:
        rgb = font_or_fill.color.rgb if hasattr(font_or_fill, "color") else font_or_fill.fore_color.rgb
        return f"#{rgb}" if rgb else None
    except Exception:
        return None


def parse(pptx: Path) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from PIL import Image
    p = Presentation(str(pptx)); W, H = p.slide_width, p.slide_height
    kx, ky = SW / W, SH / H
    box = lambda sh: [round(sh.left * kx), round(sh.top * ky), round(sh.width * kx), round(sh.height * ky)]
    pt2px = (4 / 3) * (SW / (W / 9525))                # 글자 크기 pt → 화면 px
    idx = {id(s): i + 1 for i, s in enumerate(p.slides)}
    out = []
    for n, s in enumerate(p.slides, 1):
        text = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width > W * .9 and sh.top < H * .05 and sh.height > H * .6]
        rec = {"n": n, "song": bool(pics), "creed": bool(re.search(r"사도신경|Символ веры|Creed", text)), "title_sig": ""}
        if pics:
            pic = pics[0]; blob = pic.image.blob
            im = Image.open(io.BytesIO(blob)).convert("L"); w, h = im.size
            t = im.crop((0, 0, int(w * .5), int(h * .13))).resize((32, 8)); v = list(t.get_flattened_data()) if hasattr(t, 'get_flattened_data') else list(t.getdata()); avg = sum(v) / len(v)
            rec["title_sig"] = "".join("1" if x > avg else "0" for x in v)
            sub, chips = None, []
            for sh in s.shapes:
                if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.top < H * .6:
                    continue
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.height < H * .08:      # 절 바로가기 단추
                    to = None
                    try:
                        tgt = sh.click_action.target_slide
                        to = idx.get(id(tgt)) if tgt is not None else None
                    except Exception:
                        pass
                    r0 = next((r for pa in sh.text_frame.paragraphs for r in pa.runs), None)
                    chips.append({"t": sh.text_frame.text.strip(), "to_n": to, "on": bool(r0 and _hex(r0.font) == "#FF0000"), "box": box(sh)})
                elif sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and sub is None:              # 자막 글상자
                    lines, size = [], None
                    for pa in sh.text_frame.paragraphs:
                        tx = "".join(r.text for r in pa.runs).strip()
                        if not tx: continue
                        r0 = pa.runs[0]
                        if r0.font.size and not size: size = r0.font.size.pt
                        lines.append({"t": tx, "c": _hex(r0.font) or "#ffffff"})
                    if lines:
                        fill = None
                        try:
                            fill = _hex(sh.fill) if sh.fill.type == 1 else None
                        except Exception:
                            pass
                        sub = {"box": box(sh), "size": round((size or 40) * pt2px), "fill": fill or "#000000", "lines": lines}
            rec.update({"blob": blob, "pic": box(pic), "sub": sub, "chips": chips})
        out.append(rec)
    return out


def song_record(title: str, slides: list[dict], a: int, b: int, source: str) -> dict:
    part = [s for s in slides if a <= s["n"] <= b]
    def rel(to_n): return (to_n - a) if to_n is not None and a <= to_n <= b else None
    return {"title": title, "source": source, "from": [a, b], "slides": [
        {"img": base64.b64encode(s["blob"]).decode(), "pic": s["pic"], "sub": s["sub"],
         "chips": [{"t": c["t"], "to": rel(c["to_n"]), "on": c["on"]} for c in s["chips"]]} for s in part]}


# ── 화면 ────────────────────────────────────────────────────────
def render(song: dict) -> str:
    data = json.dumps(song, ensure_ascii=False).replace("</", "<\\/")
    return PAGE.replace("@@TITLE@@", html.escape(song["title"])).replace("@@N@@", str(len(song["slides"]))).replace("@@DATA@@", data)


PAGE = r"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex"><title>@@TITLE@@ · 곡별 PPT</title>
<link rel="stylesheet" href="https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable.min.css">
<style>
*{box-sizing:border-box}html{-webkit-text-size-adjust:100%}
body{margin:0;background:#0f172a;color:#fff;font-family:'Pretendard Variable',Pretendard,'Apple SD Gothic Neo',sans-serif}
.top{position:sticky;top:0;z-index:20;display:flex;flex-wrap:wrap;gap:8px;align-items:center;padding:8px 12px;background:#1f2937}
.top b{margin-right:auto;font-size:15px}
.top button{font:700 13px inherit;border:0;border-radius:999px;padding:7px 14px;background:#fff;color:#111;cursor:pointer}
.top button.on{background:#f6c76b}.top button:disabled{opacity:.45;cursor:default}
#msg{font-size:12px;color:#fcd34d}
main{padding:12px;display:flex;flex-direction:column;align-items:center;gap:12px}
.sl{position:relative;width:min(100%,calc((100vh - 230px)*16/9));aspect-ratio:16/9;overflow:hidden;background:#fff;border-radius:6px}
.sl .in{position:absolute;left:0;top:0;width:1920px;height:1080px;transform-origin:0 0;background:#fff}
.L1{position:absolute;display:block}
.L2{position:absolute;display:flex;flex-direction:column;align-items:center;justify-content:center;text-align:center;line-height:1.18;padding:0 30px;font-family:Calibri,'Pretendard Variable',sans-serif;font-weight:600}
.L2 div{white-space:nowrap;outline:none}
body.edit .L2 div{background:rgba(255,255,255,.12);border-radius:6px;outline:2px dashed #f6c76b;margin:2px 0;padding:0 8px;min-width:200px}
/* 3층 절 바로가기 — 앞 화면 글자를 가리지 않게 슬라이드 아래 줄로 (2026-10-03) */
.L3{display:flex;flex-wrap:wrap;gap:6px;justify-content:center;max-width:1400px}
.L3 button{font:600 13px inherit;background:#1e293b;color:#e2e8f0;border:1px solid #334155;border-radius:999px;padding:6px 12px;cursor:pointer}
.L3 button.on{background:#f6c76b;color:#111;border-color:#f6c76b}
.L0{position:absolute;left:10px;top:8px;z-index:5;font:700 26px 'Pretendard Variable',sans-serif;color:#fff;background:rgba(15,23,42,.72);border-radius:999px;padding:4px 16px}
body.nolabel .L0{display:none}
#nav{display:flex;gap:10px;align-items:center;font-weight:700}
#nav button{font:700 14px inherit;border:0;border-radius:999px;padding:8px 18px;background:#fff;color:#111;cursor:pointer}
#th{display:flex;flex-wrap:wrap;gap:8px;justify-content:center;max-width:1400px}
#th .t{width:160px;cursor:pointer;border:3px solid transparent;border-radius:8px}#th .t.on{border-color:#f6c76b}
#th .t .sl{width:100%;border-radius:4px}
</style></head><body>
<div class="top"><b>🎵 @@TITLE@@ · @@N@@장</b><span id="msg"></span>
<button id="bl" onclick="document.body.classList.toggle('nolabel')">곡명 표시</button>
<button id="be" onclick="toggleEdit()">✏️ 자막 편집</button><button id="bs" onclick="save()" disabled>💾 저장</button></div>
<main><section class="sl" id="stage"><div class="in"></div></section>
<div class="L3 chips" id="chips"></div>
<div id="nav"><button onclick="show(cur-1)">◀ 이전</button><span id="pn"></span><button onclick="show(cur+1)">다음 ▶</button></div>
<div id="th"></div></main>
<script id="song" type="application/json">@@DATA@@</script>
<script>
const S=JSON.parse(document.getElementById('song').textContent), N=S.slides.length;
let cur=0, dirty=false;
function el(t,c){const e=document.createElement(t); if(c)e.className=c; return e;}
function build(i,inn,live){ inn.replaceChildren(); const s=S.slides[i];
  const lab=el('div','L0'); lab.textContent=S.title+' '+(i+1)+'/'+N; inn.appendChild(lab);            // 왼쪽 위 곡명 n/N
  const im=el('img','L1 score'); im.src='data:image/png;base64,'+s.img;                                  // 1층 악보 그림
  Object.assign(im.style,{left:s.pic[0]+'px',top:s.pic[1]+'px',width:s.pic[2]+'px',height:s.pic[3]+'px'}); inn.appendChild(im);
  if(s.sub){ const b=el('div','L2 sub'); const x=s.sub.box;                                              // 2층 자막(글자)
    Object.assign(b.style,{left:x[0]+'px',top:x[1]+'px',width:x[2]+'px',height:x[3]+'px',background:s.sub.fill,fontSize:s.sub.size+'px'});
    s.sub.lines.forEach((ln,k)=>{ const d=el('div'); d.textContent=ln.t; d.style.color=ln.c;
      if(live){ d.contentEditable=document.body.classList.contains('edit'); d.oninput=()=>{ ln.t=d.textContent; dirty=true; mark(); }; }
      b.appendChild(d); }); inn.appendChild(b); }
  if(live){ const c=document.getElementById('chips'); c.replaceChildren();                                 // 3층 절 바로가기(슬라이드 아래)
    s.chips.forEach(ch=>{ const bt=el('button',ch.on?'on':''); bt.textContent=ch.t; if(ch.to!==null) bt.onclick=()=>show(ch.to); c.appendChild(bt); }); } }
// 자막 글자 맞춤 — 원본 글꼴(Calibri)이 없는 기기에서도 칸 밖으로 넘치지 않게 줄인다
function fitSub(inn){ const b=inn.querySelector('.L2'); if(!b) return; let f=parseFloat(b.style.fontSize);
  for(let k=0;k<30&&f>18&&(b.scrollHeight>b.clientHeight+2||[...b.children].some(d=>d.scrollWidth>b.clientWidth-160));k++){ f*=0.94; b.style.fontSize=f+'px'; } }
function fit(){ document.querySelectorAll('.sl').forEach(sl=>{ sl.querySelector('.in').style.transform='scale('+(sl.clientWidth/1920)+')'; }); }
function show(i){ cur=Math.max(0,Math.min(N-1,i)); const inn=document.querySelector('#stage .in'); build(cur,inn,true); fitSub(inn);
  document.getElementById('pn').textContent=(cur+1)+' / '+N; document.querySelectorAll('#th .t').forEach((t,k)=>t.classList.toggle('on',k===cur)); fit(); }
function thumbs(){ const th=document.getElementById('th'); for(let i=0;i<N;i++){ const t=el('div','t'), sl=el('section','sl'), inn=el('div','in');
  sl.appendChild(inn); t.appendChild(sl); build(i,inn,false); t.onclick=()=>show(i); th.appendChild(t); fitSub(inn);} }
function mark(){ document.getElementById('bs').disabled=!dirty; document.getElementById('msg').textContent=dirty?'고친 자막이 있습니다 — 💾 저장':''; }
function toggleEdit(){ document.body.classList.toggle('edit'); document.getElementById('be').classList.toggle('on'); show(cur); }
const BASES=[location.protocol.startsWith('http')?location.origin:null,'http://127.0.0.1:8765','http://100.91.67.59:8765'].filter((v,i,a)=>v&&a.indexOf(v)===i);
async function save(){ const m=document.getElementById('msg'); m.textContent='저장 중…';
  const body=JSON.stringify({title:S.title, subs:S.slides.map(s=>s.sub?s.sub.lines.map(l=>l.t):null)});
  for(const b of BASES){ try{ const r=await fetch(b+'/songppt/save',{method:'POST',headers:{'Content-Type':'application/json'},body}); if(r.ok){ const j=await r.json(); dirty=false; mark(); m.textContent='✅ 저장했습니다'+(j.drive?' · 드라이브 갱신':''); return; } }catch(e){} }
  m.textContent='저장 실패 — 맥미니 저장 서버(docsave)에 닿지 않습니다. 맥미니·테일스케일에서 열어 주세요.'; }
addEventListener('keydown',e=>{ if(document.body.classList.contains('edit')&&e.target.isContentEditable) return;
  if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();show(cur+1)} else if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(cur-1)} });
addEventListener('resize',fit); thumbs(); show(0);
</script></body></html>"""


# ── 저장·올리기 ─────────────────────────────────────────────────
def write(song: dict) -> Path:
    d = DB / file_name(song["title"])[:-5]; d.mkdir(parents=True, exist_ok=True)
    (d / "song.json").write_text(json.dumps(song, ensure_ascii=False), encoding="utf-8")
    OUT.mkdir(parents=True, exist_ok=True)
    o = OUT / file_name(song["title"]); o.write_text(render(song), encoding="utf-8")
    return o


def drive():
    sys.path.insert(0, str(HERE.parent)); import prep
    return prep, prep.G(prep.access_token())


def upload(o: Path) -> dict:
    import uuid
    prep, g = drive()
    folder = g.find_child(DRIVE_ROOT, DRIVE_FOLDER_NAME, prep.FOLDER_MIME) or g.create_folder(DRIVE_ROOT, DRIVE_FOLDER_NAME)
    old = g.find_child(folder["id"], o.name)
    b = f"b{uuid.uuid4().hex}"; mime = "text/html"
    meta = json.dumps({"name": o.name} if old else {"name": o.name, "parents": [folder["id"]], "mimeType": mime}).encode()
    body = (f"--{b}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n".encode() + meta +
            f"\r\n--{b}\r\nContent-Type: {mime}\r\n\r\n".encode() + o.read_bytes() + f"\r\n--{b}--".encode())
    if old:
        url = f"{prep.UPLOAD}/files/{old['id']}?uploadType=multipart&supportsAllDrives=true&fields=id,name,webViewLink"
        return {**g.req("PATCH", url, body, ctype=f"multipart/related; boundary={b}", timeout=600), "status": "updated"}
    url = f"{prep.UPLOAD}/files?uploadType=multipart&supportsAllDrives=true&fields=id,name,webViewLink"
    return {**g.req("POST", url, body, ctype=f"multipart/related; boundary={b}", timeout=600), "status": "created"}


def save_subs(req: dict, do_upload: bool = True) -> dict:
    """화면에서 고친 자막 저장 — {title, subs:[[줄,…]|null, …]}."""
    p = DB / file_name(req["title"])[:-5] / "song.json"
    song = json.loads(p.read_text(encoding="utf-8"))
    for s, lines in zip(song["slides"], req["subs"]):
        if s["sub"] and lines is not None:
            for ln, t in zip(s["sub"]["lines"], lines):
                ln["t"] = t
    o = write(song)
    return {"ok": True, "file": str(o), "drive": upload(o).get("webViewLink") if do_upload else None}


def extract(pptx: Path, date: str, only: str | None = None, do_upload: bool = False) -> list[tuple]:
    sys.path.insert(0, str(HERE)); import slides as SL
    sl = parse(pptx)
    named = name_segments(segments(sl), SL.song_titles(date))
    done = []
    for title, a, b in named:
        if only and title != only:
            continue
        o = write(song_record(title, sl, a, b, f"{date} 주일예배 PPT {a}~{b}장"))
        link = upload(o).get("webViewLink") if do_upload and not title.startswith("확인필요") else None
        done.append((title, a, b, o, link))
    return done


if __name__ == "__main__":
    a = sys.argv[1:]
    if a and a[0] == "save":
        print(json.dumps(save_subs(json.load(sys.stdin)), ensure_ascii=False)); sys.exit(0)
    only = a[a.index("--only") + 1] if "--only" in a else None
    for t, s, e, o, link in extract(Path(a[1]), a[2], only, "--upload" in a):
        print(f"{t}\t{s}~{e}장\t{o}" + (f"\t{link}" if link else ""))
